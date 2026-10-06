from __future__ import annotations

import csv
from io import BytesIO, StringIO

from fastapi.responses import Response, StreamingResponse
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet


def _stream_bulk_export_csv(filename: str, rows: list[dict[str, object]]) -> StreamingResponse:
    content = _render_bulk_export_csv_bytes(rows)
    return StreamingResponse(
        iter([content]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _render_bulk_export_csv_bytes(rows: list[dict[str, object]]) -> bytes:
    headers = list(rows[0].keys()) if rows else []
    buffer = StringIO()
    writer = csv.DictWriter(buffer, fieldnames=headers)
    if headers:
        writer.writeheader()
        for row in rows:
            writer.writerow(row)
    return buffer.getvalue().encode("utf-8")


def _stream_bulk_export_xlsx(filename: str, rows: list[dict[str, object]]) -> Response:
    content = _render_bulk_export_xlsx_bytes(rows)
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Content-Length": str(len(content)),
        },
    )


def _write_bulk_export_hyperlink(
    sheet: Worksheet,
    row_idx: int,
    link_col: int | None,
    apri_col: int | None,
) -> None:
    if None in (link_col, apri_col):
        return
    link_value = sheet.cell(row=row_idx, column=link_col).value
    if not link_value:
        return
    link_cell = sheet.cell(row=row_idx, column=link_col).coordinate
    sheet.cell(row=row_idx, column=apri_col).value = f'=HYPERLINK({link_cell},"Clicca qui")'


def _render_bulk_export_xlsx_bytes(rows: list[dict[str, object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "intestatari"
    headers = list(rows[0].keys()) if rows else []
    if headers:
        sheet.append(headers)
        link_col = headers.index("link_involture") + 1 if "link_involture" in headers else None
        apri_col = headers.index("apri_involture") + 1 if "apri_involture" in headers else None
        for row_idx, row in enumerate(rows, start=2):
            sheet.append([row.get(header, "") for header in headers])
            _write_bulk_export_hyperlink(sheet, row_idx, link_col, apri_col)
    buffer = BytesIO()
    workbook.save(buffer)
    content = buffer.getvalue()
    workbook.close()
    return content
