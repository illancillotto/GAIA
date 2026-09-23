from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path

from codicefiscale import control_code, isvalid
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.modules.presenze.models import PresenzeCollaborator


@dataclass(frozen=True)
class TaxCodeImportEntry:
    employee_code: str
    tax_code: str
    source: str


class TaxCodeImportError(ValueError):
    """Raised when an import cannot be applied without risking partial updates."""

    def __init__(self, message: str, report: dict[str, int]) -> None:
        super().__init__(message)
        self.report = report


def normalize_tax_code(value: object) -> str | None:
    if value in (None, ""):
        return None
    text = "".join(str(value).strip().upper().split())
    if not text or not isvalid(text) or control_code(text[:15]) != text[15]:
        raise ValueError(f"Codice fiscale non valido: {value!r}.")
    return text


def _tax_code_from_record_key(value: object) -> str | None:
    if value in (None, ""):
        return None
    text = str(value).strip()
    suffix = text.split("-", 1)[1].strip() if "-" in text else text
    return None if suffix.isdigit() else normalize_tax_code(suffix)


def _employee_code(value: object) -> str:
    return str(value).strip() if value not in (None, "") else ""


def _add_candidate(
    candidates: dict[str, TaxCodeImportEntry],
    employee_code: str,
    tax_code: str,
    source: str,
) -> None:
    previous = candidates.get(employee_code)
    if previous is not None and previous.tax_code != tax_code:
        raise ValueError(f"Codici fiscali discordanti per matricola {employee_code}.")
    candidates.setdefault(employee_code, TaxCodeImportEntry(employee_code, tax_code, source))


def _extract_operai_tax_codes(workbook) -> dict[str, TaxCodeImportEntry]:
    if "Operai" not in workbook.sheetnames:
        return {}
    candidates: dict[str, TaxCodeImportEntry] = {}
    for row_number, row in enumerate(
        workbook["Operai"].iter_rows(min_row=2, values_only=True), start=2
    ):
        employee_code = _employee_code(row[3] if len(row) > 3 else None)
        tax_code = normalize_tax_code(row[15] if len(row) > 15 else None)
        if employee_code and tax_code:
            _add_candidate(candidates, employee_code, tax_code, f"Operai!P{row_number}")
    return candidates


def _extract_riepilogo_tax_codes(workbook) -> dict[str, TaxCodeImportEntry]:
    if "Riepilogo" not in workbook.sheetnames:
        return {}
    candidates: dict[str, TaxCodeImportEntry] = {}
    for row_number, row in enumerate(
        workbook["Riepilogo"].iter_rows(min_row=2, values_only=True), start=2
    ):
        employee_code = _employee_code(row[0] if row else None)
        tax_code = _tax_code_from_record_key(row[2] if len(row) > 2 else None)
        if employee_code and tax_code:
            _add_candidate(candidates, employee_code, tax_code, f"Riepilogo!C{row_number}")
    return candidates


def extract_tax_codes(workbook_path: Path) -> list[TaxCodeImportEntry]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        candidates = _extract_operai_tax_codes(workbook)
        for employee_code, entry in _extract_riepilogo_tax_codes(workbook).items():
            _add_candidate(candidates, employee_code, entry.tax_code, entry.source)
        return sorted(
            candidates.values(),
            key=lambda item: (0, int(item.employee_code))
            if item.employee_code.isdigit()
            else (1, item.employee_code),
        )
    finally:
        workbook.close()


def _validated_entries(entries: list[TaxCodeImportEntry]) -> list[TaxCodeImportEntry]:
    candidates: dict[str, TaxCodeImportEntry] = {}
    for entry in entries:
        tax_code = normalize_tax_code(entry.tax_code)
        if not entry.employee_code.strip() or tax_code is None:
            raise ValueError("Matricola e codice fiscale sono obbligatori.")
        _add_candidate(candidates, entry.employee_code.strip(), tax_code, entry.source)
    return list(candidates.values())


def _plan_import(collaborators, entries):
    by_employee: dict[str, list[PresenzeCollaborator]] = {}
    for collaborator in collaborators:
        by_employee.setdefault(collaborator.employee_code, []).append(collaborator)
    report = dict(
        entries=len(entries), matched=0, missing=0, conflicts=0, changed=0, unchanged=0, applied=0
    )
    changes: list[tuple[PresenzeCollaborator, str]] = []
    for entry in entries:
        rows = by_employee.get(entry.employee_code, [])
        if len(rows) != 1:
            report["conflicts" if rows else "missing"] += 1
            continue
        report["matched"] += 1
        collaborator = rows[0]
        existing = normalize_tax_code(collaborator.tax_code)
        if existing == entry.tax_code:
            report["unchanged"] += 1
        elif existing is not None:
            report["conflicts"] += 1
        else:
            report["changed"] += 1
            changes.append((collaborator, entry.tax_code))
    return report, changes


def import_tax_codes(
    db: Session,
    entries: list[TaxCodeImportEntry],
    *,
    apply: bool,
    company_code: str | None = None,
    before_commit: Callable | None = None,
) -> dict[str, int]:
    entries = _validated_entries(entries)
    statement = select(PresenzeCollaborator).where(
        PresenzeCollaborator.employee_code.in_([entry.employee_code for entry in entries])
    )
    if company_code is not None:
        statement = statement.where(PresenzeCollaborator.company_code == company_code)
    if apply:
        statement = statement.with_for_update()
    report, changes = _plan_import(db.scalars(statement).all(), entries)
    if not apply:
        return report
    try:
        if report["missing"] or report["conflicts"]:
            raise TaxCodeImportError("Importazione bloccata: dati mancanti o discordanti.", report)
        if before_commit is not None:
            before_commit(report, changes)
        for collaborator, tax_code in changes:
            collaborator.tax_code = tax_code
        db.commit()
        report["applied"] = report["changed"]
    except Exception:
        db.rollback()
        raise
    return report
