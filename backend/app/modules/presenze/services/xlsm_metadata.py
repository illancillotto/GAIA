from __future__ import annotations

from dataclasses import dataclass, replace

from openpyxl.worksheet.worksheet import Worksheet

from app.modules.presenze.services.tax_code_import import _extract_riepilogo_tax_codes


@dataclass(frozen=True)
class OperaiMetadata:
    tax_code: str | None
    qualifica: str | None
    mansione: str | None
    inquadramento: str | None
    period_text: str | None


def format_operai_date(value: object | None) -> str:
    if value is None:
        return "--"
    if hasattr(value, "strftime"):
        return value.strftime("%d-%m-%y")
    return str(value).strip() or "--"


def build_operai_period_text(
    dal: object | None,
    al: object | None,
    proroga: object | None,
    riass1_dal: object | None,
    riass1_al: object | None,
) -> str | None:
    values = [dal, al, proroga, riass1_dal, riass1_al]
    if all(value in (None, "", "--") for value in values):
        return None
    return (
        f"Dal {format_operai_date(dal)} al {format_operai_date(al)}"
        f"        Proroga al {format_operai_date(proroga)}"
        f"                            Riass.dal {format_operai_date(riass1_dal)} al {format_operai_date(riass1_al)}"
    )


def load_operai_metadata(ws: Worksheet) -> dict[int, OperaiMetadata]:
    metadata: dict[int, OperaiMetadata] = {}
    for row in range(2, ws.max_row + 1):
        employee_code = ws.cell(row, 4).value
        if not isinstance(employee_code, int):
            continue
        metadata[employee_code] = OperaiMetadata(
            tax_code=str(ws.cell(row, 16).value).strip()
            if ws.cell(row, 16).value not in (None, "")
            else None,
            qualifica=str(ws.cell(row, 5).value).strip()
            if ws.cell(row, 5).value not in (None, "")
            else None,
            mansione=str(ws.cell(row, 6).value).strip()
            if ws.cell(row, 6).value not in (None, "")
            else None,
            inquadramento=str(ws.cell(row, 7).value).strip()
            if ws.cell(row, 7).value not in (None, "")
            else None,
            period_text=build_operai_period_text(
                ws.cell(row, 8).value,
                ws.cell(row, 9).value,
                ws.cell(row, 10).value,
                ws.cell(row, 11).value,
                ws.cell(row, 12).value,
            ),
        )
    return metadata


def load_workbook_metadata(workbook) -> dict[int, OperaiMetadata]:
    operai = workbook["Operai"] if "Operai" in workbook.sheetnames else None
    metadata = load_operai_metadata(operai) if operai is not None else {}
    for code, entry in _extract_riepilogo_tax_codes(workbook).items():
        employee = int(code)
        current = metadata.get(employee, OperaiMetadata(None, None, None, None, None))
        if not current.tax_code:
            metadata[employee] = replace(current, tax_code=entry.tax_code)
    return metadata


def resolve_archive_record_source(source_value, collaborator, metadata) -> object:
    """Prefer GAIA's saved code, then workbook metadata, then the historical key."""
    return collaborator.tax_code or (metadata.tax_code if metadata else None) or source_value
