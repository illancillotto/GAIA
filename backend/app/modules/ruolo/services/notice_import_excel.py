"""Bounded parser for the approved 2022/2023 workbook, never a notification rule."""

import hashlib
import json
from collections import Counter
from datetime import date, datetime
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MAX_BYTES = 12 * 1024 * 1024
MAX_ROWS = 10000
PARSER_VERSION = "registro-2022-2023-v1"
HEADERS = {
    2: "Rif. 2022",
    3: "Rif. 2023",
    7: "MOTIVAZIONE",
    8: "DATA NOTIFICA",
    20: "Cod.Fiscale",
    21: "Avviso n.",
    56: "Anno",
}


def json_bytes(value: object) -> bytes:
    return json.dumps(
        value, ensure_ascii=True, sort_keys=True, separators=(",", ":"), default=str
    ).encode()


def digest(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def cell_value(value):
    if isinstance(value, datetime | date):
        return value.isoformat()
    return value


def reference(value) -> str | None:
    if value is None or str(value).strip() in {"", "---", "-"}:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalized_date(value) -> tuple[str | None, str | None]:
    raw = reference(value)
    if raw is None:
        return None, None
    correction = raw == "29/06/204"
    if correction:
        raw = "29/06/2024"
    try:
        parsed = (
            datetime.strptime(raw, "%d/%m/%Y").date()
            if "/" in raw
            else date.fromisoformat(raw[:10])
        )
    except ValueError:
        return None, "data_non_valida"
    if parsed > date.today() or parsed.year < 1900:
        return None, "data_non_valida"
    return parsed.isoformat(), "data_corretta_29_06_204" if correction else None


def annual_positions(values: list, anomalies: list) -> list[dict]:
    positions = []
    for index, year in ((2, 2022), (3, 2023)):
        ref = reference(values[index])
        if ref is None:
            continue
        if len(ref) > 100:
            anomalies.append(f"riferimento_{year}_non_valido")
            continue
        if not isinstance(values[index], str):
            anomalies.append(f"riferimento_{year}_numerico_verificare_zeri")
        positions.append({"source_namespace": "incass", "source_reference": ref, "tax_year": year})
    if not positions:
        anomalies.append("posizioni_assenti")
    return positions


def parse_row(values: list, row_number: int) -> dict:
    anomalies = ["notifica_da_verificare", "step_da_verificare"]
    dates = {}
    for index in range(8, 12):
        column = get_column_letter(index + 1)
        normalized, warning = normalized_date(values[index])
        dates[column] = normalized
        if warning:
            anomalies.append(f"{column}:{warning}")
    if len({value for value in dates.values() if value}) > 1:
        anomalies.append("date_discordanti")
    number = reference(values[21])
    tax_code = reference(values[20])
    if tax_code is None or len(tax_code) > 20:
        anomalies.append("codice_fiscale_assente_o_non_valido")
        tax_code = None
    original = {get_column_letter(i + 1): cell_value(value) for i, value in enumerate(values)}
    fingerprint = digest(json_bytes(original))
    if number is None or len(number) > 100:
        anomalies.append("numero_documento_assente_o_non_valido")
        number = f"Senza numero {fingerprint[:16]}"
    payload = {
        "document_number": number,
        "tax_code": tax_code,
        "positions": annual_positions(values, anomalies),
        "original": original,
        "dates": dates,
        "sheet": "Dati",
        "motivation": reference(values[7]),
    }
    return {
        "row_number": row_number,
        "source_key": number,
        "fingerprint": fingerprint,
        "payload": payload,
        "anomalies": anomalies,
    }


def _validate_archive(content: bytes):
    if not content or len(content) > MAX_BYTES:
        raise ValueError("File vuoto o oltre il limite di 12 MiB")
    with ZipFile(BytesIO(content)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 64 * 1024 * 1024:
            raise ValueError("Cartella Excel decompressa oltre il limite di 64 MiB")


def _read_sheet(sheet) -> tuple[list[dict], dict]:
    if sheet.max_row is None or sheet.max_column is None:
        raise ValueError("Dimensioni del foglio Dati assenti")
    if sheet.max_row > MAX_ROWS or sheet.max_column > 128:
        raise ValueError("Foglio Dati oltre i limiti supportati")
    values = sheet.iter_rows(values_only=True)
    headers = next(values)
    if any(len(headers) <= index or headers[index] != label for index, label in HEADERS.items()):
        raise ValueError("Tracciato diverso dal file avvisi 2022/2023 approvato")
    rows, ignored = [], 0
    for row_number, row in enumerate(values, 2):
        if reference(row[56]) not in {"2022", "2023"}:
            ignored += 1
            continue
        rows.append(parse_row(list(row), row_number))
    if not rows:
        raise ValueError("Nessuna riga 2022/2023 nel foglio Dati")
    return rows, _row_summary(rows, ignored)


def _row_summary(rows: list[dict], ignored: int) -> dict:
    references = Counter(
        (position["tax_year"], position["source_reference"])
        for row in rows
        for position in row["payload"]["positions"]
    )
    for row in rows:
        if any(
            references[(p["tax_year"], p["source_reference"])] > 1
            for p in row["payload"]["positions"]
        ):
            row["anomalies"].append("riferimento_annuale_in_piu_documenti")
    return {
        "rows": len(rows),
        "ignored_non_operational": ignored,
        "annual_references": sum(references.values()),
        "distinct_annual_references": len(references),
    }


def parse_excel(content: bytes) -> tuple[list[dict], dict]:
    try:
        _validate_archive(content)
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
    except (BadZipFile, KeyError, OSError, SyntaxError) as exc:
        raise ValueError("Cartella XLSX non valida") from exc
    try:
        if "Dati" not in workbook.sheetnames:
            raise ValueError("Foglio Dati assente")
        return _read_sheet(workbook["Dati"])
    except SyntaxError as exc:
        raise ValueError("XML del foglio Dati non valido") from exc
    finally:
        workbook.close()
