"""Validate the source workbook before writing any personal access data."""

from collections import Counter
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.modules.ruolo.portal_access_models import RuoloTaxpayerPortalAccess
from app.services.elaborazioni_credentials import get_credential_fernet


@dataclass(frozen=True)
class PortalAccessRow:
    tax_code: str
    account_code: str
    pin: str = field(repr=False)


def _digits(cell) -> str:
    value = cell.value
    if isinstance(value, str):
        value = value.strip()
    if isinstance(value, bool):
        raise ValueError("Credenziale non valida")
    if isinstance(value, int):
        value = str(value)
        if set(cell.number_format) == {"0"}:
            value = value.zfill(len(cell.number_format))
    if not isinstance(value, str) or not value.isascii() or not value.isdigit():
        raise ValueError("Credenziale non valida")
    if len(value) > 32:
        raise ValueError("Credenziale troppo lunga")
    return value


def _access_row(cells, columns: list[int], number: int) -> PortalAccessRow | None:
    cf_col, account_col, pin_col = columns
    tax_code = str(cells[cf_col].value or "").strip().upper()
    if not tax_code and cells[account_col].value is None and cells[pin_col].value is None:
        return None
    if not tax_code.isascii() or not tax_code.isalnum() or len(tax_code) not in (11, 16):
        raise ValueError(f"Codice fiscale non valido alla riga {number}")
    try:
        return PortalAccessRow(tax_code, _digits(cells[account_col]), _digits(cells[pin_col]))
    except ValueError as exc:
        raise ValueError(f"Credenziali non valide alla riga {number}") from exc


def _collect_access_rows(rows, columns: list[int]):
    accounts: dict[str, PortalAccessRow] = {}
    counts: Counter = Counter()
    for number, cells in enumerate(rows, 2):
        row = _access_row(cells, columns, number)
        if row is None:
            counts["summary_or_blank_rows"] += 1
            continue
        if row.account_code in accounts:
            if accounts[row.account_code] != row:
                raise ValueError(f"Codice ditta in conflitto alla riga {number}")
            counts["duplicate_rows"] += 1
        accounts[row.account_code] = row
        counts["source_rows"] += 1
    if not accounts:
        raise ValueError("Nessuna credenziale nel file")
    counts["accounts"] = len(accounts)
    counts["tax_codes"] = len({row.tax_code for row in accounts.values()})
    return list(accounts.values()), dict(counts)


def read_access_workbook(path: Path) -> tuple[list[PortalAccessRow], dict[str, int]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        rows = workbook.active.iter_rows()
        headers = [cell.value for cell in next(rows, [])]
        required = ("CodFisc", "Utenza", "PIN")
        if any(headers.count(name) != 1 for name in required):
            raise ValueError("Colonne richieste: CodFisc, Utenza, PIN (univoche)")
        return _collect_access_rows(rows, [headers.index(name) for name in required])
    finally:
        workbook.close()


def import_access_rows(
    db: Session, rows: list[PortalAccessRow], source_sha256: str
) -> dict[str, int]:
    """Upsert in the caller transaction; never reassign an existing account to another CF."""
    fernet = get_credential_fernet()
    existing = {row.account_code: row for row in db.scalars(select(RuoloTaxpayerPortalAccess))}
    if any(
        row.account_code in existing and existing[row.account_code].tax_code != row.tax_code
        for row in rows
    ):
        raise ValueError("Import interrotto: codice ditta già associato a un altro codice fiscale")
    counts = {"created": 0, "updated": 0, "unchanged": 0}
    for row in rows:
        stored = existing.get(row.account_code)
        if stored is not None and fernet.decrypt(stored.pin_encrypted).decode() == row.pin:
            counts["unchanged"] += 1
            continue
        counts["created" if stored is None else "updated"] += 1
        if stored is None:
            stored = RuoloTaxpayerPortalAccess(account_code=row.account_code, tax_code=row.tax_code)
            db.add(stored)
            existing[row.account_code] = stored
        stored.pin_encrypted = fernet.encrypt(row.pin.encode())
        stored.source_sha256 = source_sha256
        stored.imported_at = datetime.now(UTC)
    db.flush()
    return counts
