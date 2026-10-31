"""Serialize Capacitas grid rows (the data behind a Kendo "Esporta") into an xlsx workbook."""

from __future__ import annotations

import json
import re
from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook

_SCALAR_TYPES = (int, float, bool)
_MAX_SHEET_TITLE = 31
_ASPNET_DATE_RE = re.compile(r"^/Date\((-?\d+)(?:[+-]\d{4})?\)/$")
_PORTAL_TIMEZONE = ZoneInfo("Europe/Rome")


@dataclass(frozen=True, slots=True)
class ExportPayload:
    """One exported workbook and the number of data rows it holds (header excluded)."""

    content: bytes
    rows: int


def _has_content(cell: object) -> bool:
    # The inCASS export ends with a row holding only a newline in the "Reso" column.
    return cell.strip() != "" if isinstance(cell, str) else cell is not None


def xlsx_data_rows(content: bytes) -> int:
    # Read-only sheets may not declare their dimension, so count the non-empty rows.
    workbook = load_workbook(BytesIO(content), read_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        return max(sum(1 for row in rows if any(_has_content(cell) for cell in row)) - 1, 0)
    finally:
        workbook.close()


def collect_columns(rows: Iterable[Mapping[str, object]]) -> list[str]:
    """Return every key seen in the rows, preserving first-seen order."""
    columns: dict[str, None] = {}
    for row in rows:
        columns.update(dict.fromkeys(row))
    return list(columns)


def _string_value(value: str) -> object:
    match = _ASPNET_DATE_RE.match(value)
    if match is None:
        return value
    instant = datetime.fromtimestamp(int(match.group(1)) / 1000, tz=UTC)
    # Excel has no time zone: store the wall-clock time shown by the portal.
    return instant.astimezone(_PORTAL_TIMEZONE).replace(tzinfo=None)


def cell_value(value: object) -> object:
    if isinstance(value, str):
        return _string_value(value)
    if value is None or isinstance(value, _SCALAR_TYPES):
        return value
    return json.dumps(value, ensure_ascii=False, default=str)


def rows_to_payload(rows: list[Mapping[str, object]], *, sheet_title: str) -> ExportPayload:
    return ExportPayload(content=rows_to_xlsx(rows, sheet_title=sheet_title), rows=len(rows))


def rows_to_xlsx(rows: list[Mapping[str, object]], *, sheet_title: str) -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title=sheet_title[:_MAX_SHEET_TITLE])
    columns = collect_columns(rows)
    sheet.append(columns)
    for row in rows:
        sheet.append([cell_value(row.get(column)) for column in columns])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
