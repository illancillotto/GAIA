"""Shift entitlement, dated authority and both GATE transports."""

import uuid
from datetime import date, datetime, time, timedelta
from types import SimpleNamespace

import pytest
from pydantic import ValidationError
from sqlalchemy import select
from test_presenze_api import (  # noqa: F401 - shared isolated API database fixture
    TestingSessionLocal,
    _create_user,
    _login,
    client,
    setup_database,
)
from test_presenze_meal_vouchers import _record

from app.core.datetime_compat import UTC
from app.modules.presenze.models import PresenzeDailyPunch, PresenzeDailyRecord
from app.modules.presenze.services.meal_vouchers import meal_voucher_values
from app.modules.presenze.services.shift_assignments import (
    ShiftAssignmentConflict,
    ShiftAssignmentOrigin,
    record_shift_assignment,
    shift_assignment_values,
)
from app.modules.presenze.services.shift_worker_rules import (
    shift_classification,
    shift_meal_voucher,
    shift_punch_minutes,
    shift_quality,
)
from app.modules.presenze.shift_worker_models import PresenzeShiftAssignment
from app.modules.presenze.shift_worker_schemas import (
    GateShiftWorkerAssignmentRequest,
    ShiftWorkerAssignmentRequest,
)
from app.services.gate_mobile_sync import (
    _apply_presenze_pending_action,
    build_presenze_giornaliere_push_payload,
)


def punch(entry, exit):
    return SimpleNamespace(
        entry_time=time.fromisoformat(entry) if entry else None,
        exit_time=time.fromisoformat(exit) if exit else None,
    )


def simple(kind="acquaiolo", punches=None, **extra):
    return SimpleNamespace(
        shift_worker_type=kind,
        shift_punches=punches or [],
        schedule_code="OPE0613",
        justified_minutes=0,
        absence_minutes=0,
        work_date=date(2026, 10, 1),
        meal_voucher_manual=False,
        raw_payload_json={},
        request_description=None,
        request_status=None,
        resolved_absence_cause="permesso",
        **extra,
    )


@pytest.mark.parametrize("kind", ["acquaiolo", "telecontrollo"])
@pytest.mark.parametrize(
    "entry,exit,night", [("06:00", "13:00", 0), ("14:00", "21:00", 0), ("22:00", "05:00", 420)]
)
def test_shift_entitlement_and_export(kind, entry, exit, night):
    punches = [punch(entry, exit)]
    record = simple(kind, punches)
    q = shift_quality(record, punches)
    assert (q.expected_minutes, q.worked_minutes, q.missing_minutes, q.status) == (
        420,
        420,
        0,
        "ok",
    )
    assert meal_voucher_values(record, 0)["meal_voucher_sources"] == ["shift"]
    record.meal_voucher_manual = True
    assert meal_voucher_values(record, 120)["meal_voucher_count"] == 1
    assert meal_voucher_values(record, 120)["meal_voucher_sources"] == ["shift", "manual"]
    for festive in (False, True):
        c = shift_classification(
            record, punches, special_day=festive, holiday_kind=None, grants_recovery_day=False
        )
        assert (c.ordinary_minutes, c.extra_minutes) == (420, 0)
        assert c.shift_festive_night_minutes == (night if festive else 0)
        assert c.shift_night_minutes == (0 if festive else night)


@pytest.mark.parametrize(
    "punches",
    [
        [],
        [punch(None, "13:00")],
        [punch("06:00", None)],
        [punch("06:00", "06:00")],
        [punch("06:00", "13:00"), punch("07:00", "08:00")],
    ],
)
def test_invalid_punches_never_grant_automatic_voucher(punches):
    record = simple(punches=punches)
    assert shift_punch_minutes(punches) is None
    assert shift_meal_voucher(record) is False
    assert meal_voucher_values(record, 600)["meal_voucher_count"] == 0
    assert shift_quality(record, punches).status == "blocking"
    assert (
        shift_classification(
            record, punches, special_day=False, holiday_kind=None, grants_recovery_day=False
        ).ordinary_minutes
        is None
    )
    record.absence_minutes = 420
    if punches:
        assert shift_quality(record, punches).status == "blocking"


def test_short_extra_rest_and_partial_absence():
    rows = [punch("06:00", "12:00")]
    record = simple(punches=rows)
    assert shift_quality(record, rows).missing_minutes == 60
    assert not shift_meal_voucher(record)
    record.justified_minutes = 60
    assert shift_quality(record, rows).missing_minutes == 0
    c = shift_classification(
        record,
        [punch("21:00", "05:00")],
        special_day=True,
        holiday_kind=None,
        grants_recovery_day=False,
    )
    assert (c.extra_minutes, c.overtime_festive_night_minutes, c.shift_festive_day_minutes) == (
        60,
        60,
        60,
    )
    record.schedule_code = "RIPTURN"
    assert shift_quality(record, []).expected_minutes == 0
    assert shift_quality(record, []).status == "ok"
    assert (
        shift_classification(
            record, [], special_day=False, holiday_kind=None, grants_recovery_day=False
        ).ordinary_minutes
        == 0
    )
    assert shift_meal_voucher(simple("none", rows)) is False


@pytest.mark.parametrize(
    "data",
    [
        {"date_from": "2026-10-20", "date_to": "2026-10-01"},
        {"date_to": "2026-11-01"},
        {"shift_worker_type": "other"},
        {"date_from": "2026-10-32"},
    ],
)
def test_assignment_validation(data):
    with pytest.raises(ValidationError):
        ShiftWorkerAssignmentRequest.model_validate(
            {
                "date_from": "2026-10-01",
                "date_to": "2026-10-31",
                "shift_worker_type": "acquaiolo",
                **data,
            }
        )
    with pytest.raises(ValidationError):
        GateShiftWorkerAssignmentRequest(
            shift_worker_type="acquaiolo",
            date_from=date(2026, 10, 1),
            date_to=date(2026, 10, 31),
            gate_shift_requested_at=datetime(2026, 10, 1),
        )


@pytest.mark.parametrize("transport", ["lan", "outbound"])
def test_gate_dated_assignment_persists_and_overrides_gaia(transport):
    actor = _create_user("shift-" + transport)
    rid, cid = _record(actor.id)
    headers = {"Authorization": "Bearer " + _login(actor.username)}
    request = {"shift_worker_type": "acquaiolo", "date_from": "2026-10-01", "date_to": "2026-10-31"}
    response = client.post(f"/presenze/giornaliere/{rid}/turnista", headers=headers, json=request)
    assert response.status_code == 200, response.text
    assert response.json()["shift_worker_source"] == "gaia"
    gate = {
        **request,
        "shift_worker_type": "telecontrollo",
        "gate_shift_requested_at": "2026-10-02T12:00:00Z",
        "gate_shift_command_id": str(uuid.uuid4()),
    }
    for _ in range(2):
        if transport == "lan":
            response = client.post(
                f"/gate/presenze/giornaliere/{rid}/patch", headers=headers, json=gate
            )
            assert response.status_code == 200, response.text
        else:
            with TestingSessionLocal() as db:
                _apply_presenze_pending_action(
                    db,
                    {
                        "id": str(uuid.uuid4()),
                        "action_type": "patch_daily_record",
                        "payload": {**gate, "record_id": str(rid), "gaia_user_id": actor.id},
                    },
                )
    assert (
        client.post(
            f"/presenze/giornaliere/{rid}/turnista", headers=headers, json=request
        ).status_code
        == 409
    )
    with TestingSessionLocal() as db:
        assert len(db.scalars(select(PresenzeShiftAssignment)).all()) == 2
        record = db.get(PresenzeDailyRecord, rid)
        db.add(
            PresenzeDailyPunch(
                daily_record_id=rid, sequence=1, entry_time=time(14), exit_time=time(21)
            )
        )
        future = PresenzeDailyRecord(
            collaborator_id=cid,
            owner_user_id=actor.id,
            work_date=date(2026, 10, 3),
            schedule_code="OPE0613",
            reperibilita_unit="none",
            validation_status="pending",
        )
        db.add(future)
        db.flush()
        assert shift_assignment_values(future)["shift_worker_type"] == "telecontrollo"
        payload = build_presenze_giornaliere_push_payload(db, month="2026-10")
        day = next(r for r in payload["giornaliere"] if r["record_id"] == str(rid))
        assert day["shift_worker_type"] == "telecontrollo"
        assert day["meal_voucher_count"] == 1
        assert day["export_ordinary_minutes"] == 420
        stamp = datetime(2026, 10, 2, 12, tzinfo=UTC)
        removal = ShiftWorkerAssignmentRequest(
            shift_worker_type="none", date_from=future.work_date, date_to=future.work_date
        )
        record_shift_assignment(
            db,
            record,
            removal,
            origin=ShiftAssignmentOrigin(
                actor.id, "gate", stamp + timedelta(seconds=1), str(uuid.uuid4())
            ),
        )
        assert shift_assignment_values(future)["shift_worker_type"] == "none"
        assert shift_assignment_values(record)["shift_worker_type"] == "telecontrollo"
        with pytest.raises(ShiftAssignmentConflict):
            record_shift_assignment(
                db,
                record,
                removal,
                origin=ShiftAssignmentOrigin(actor.id, "gaia", stamp, str(uuid.uuid4())),
            )


def test_shift_route_security_and_ranges():
    owner = _create_user("shift-owner")
    rid, _ = _record(owner.id)
    data = {"shift_worker_type": "acquaiolo", "date_from": "2026-10-01", "date_to": "2026-10-31"}
    assert client.post(f"/presenze/giornaliere/{rid}/turnista", json=data).status_code == 401
    reader = _create_user("shift-reader", role="viewer")
    headers = {"Authorization": "Bearer " + _login(reader.username)}
    assert client.post(
        f"/presenze/giornaliere/{rid}/turnista", headers=headers, json=data
    ).status_code in (403, 404)
    headers = {"Authorization": "Bearer " + _login(owner.username)}
    assert (
        client.post(
            f"/presenze/giornaliere/{rid}/turnista",
            headers=headers,
            json={**data, "date_to": "2026-11-01"},
        ).status_code
        == 422
    )


def test_gate_shift_command_requires_id():
    from app.modules.presenze.gate_daily_record_schemas import GatePresenzeDailyRecordPatchRequest

    with pytest.raises(ValidationError, match="command_id"):
        GatePresenzeDailyRecordPatchRequest(
            shift_worker_type="acquaiolo",
            date_from=date(2026, 10, 1),
            date_to=date(2026, 10, 31),
            gate_shift_requested_at=datetime.now(UTC),
        )


def test_route_enforces_edit_permission(monkeypatch):
    from app.modules.presenze.router.routes import shift_workers

    user = _create_user("shift-noedit")
    rid, _ = _record(user.id)
    monkeypatch.setattr(shift_workers, "_can_edit_daily_record", lambda *_: False)
    response = client.post(
        f"/presenze/giornaliere/{rid}/turnista",
        headers={"Authorization": "Bearer " + _login(user.username)},
        json={"shift_worker_type": "acquaiolo", "date_from": "2026-10-01", "date_to": "2026-10-31"},
    )
    assert response.status_code == 403
    with TestingSessionLocal() as db:
        assert db.scalars(select(PresenzeShiftAssignment)).all() == []


def test_shift_migration_round_trip(monkeypatch):
    import importlib.util
    from pathlib import Path

    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    from sqlalchemy import inspect
    from test_presenze_api import engine

    path = (
        Path(__file__).parents[1] / "alembic/versions/20261003_1200_presenze_shift_assignments.py"
    )
    spec = importlib.util.spec_from_file_location("shift_migration", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    with engine.begin() as connection:
        PresenzeShiftAssignment.__table__.drop(connection)
        monkeypatch.setattr(module, "op", Operations(MigrationContext.configure(connection)))
        module.upgrade()
        assert "presenze_shift_assignments" in inspect(connection).get_table_names()
        assert len(inspect(connection).get_check_constraints("presenze_shift_assignments")) == 3
        module.downgrade()
        assert "presenze_shift_assignments" not in inspect(connection).get_table_names()
        module.upgrade()


@pytest.mark.parametrize(
    "extra", [{"date_from": "2026-10-01"}, {"gate_shift_command_id": "orphan"}]
)
def test_gate_rejects_orphan_assignment_metadata(extra):
    from app.modules.presenze.gate_daily_record_schemas import GatePresenzeDailyRecordPatchRequest

    with pytest.raises(ValidationError):
        GatePresenzeDailyRecordPatchRequest.model_validate(extra)


def test_gate_rejects_combined_assignment_and_manual_values():
    from app.modules.presenze.gate_daily_record_schemas import GatePresenzeDailyRecordPatchRequest

    with pytest.raises(ValidationError, match="separatamente"):
        GatePresenzeDailyRecordPatchRequest(
            shift_worker_type="acquaiolo",
            date_from=date(2026, 10, 1),
            date_to=date(2026, 10, 31),
            gate_shift_requested_at=datetime.now(UTC),
            gate_shift_command_id="combined",
            meal_voucher_manual=True,
        )
    with pytest.raises(ValidationError, match="command_id"):
        GatePresenzeDailyRecordPatchRequest(
            shift_worker_type="acquaiolo",
            date_from=date(2026, 10, 1),
            date_to=date(2026, 10, 31),
            gate_shift_requested_at=datetime.now(UTC),
            gate_shift_command_id="",
        )


def test_split_overnight_shift_and_midnight_overlap():
    rows = [punch("22:00", "01:00"), punch("02:00", "06:00")]
    assert shift_punch_minutes(rows) == 420
    c = shift_classification(
        simple(punches=rows), rows, special_day=False, holiday_kind=None, grants_recovery_day=False
    )
    assert c.shift_night_minutes == 420
    overlapping = [punch("22:00", "01:00"), punch("00:00", "06:00")]
    assert shift_punch_minutes(overlapping) is None
    assert not shift_meal_voucher(simple(punches=overlapping))


def test_shift_voucher_exports_and_incomplete_shift_never_uses_overtime_fallback():
    from openpyxl import Workbook

    from app.modules.presenze.models import PresenzeCollaborator
    from app.modules.presenze.services.xlsm_export import (
        ExportTimesheetRow,
        write_archive2_daily_values,
        write_archivio_summary_values,
    )

    user = _create_user("shift-export")
    rid, cid = _record(user.id, extra=300)
    with TestingSessionLocal() as db:
        record = db.get(PresenzeDailyRecord, rid)
        origin = ShiftAssignmentOrigin(user.id, "gate", datetime.now(UTC), str(uuid.uuid4()))
        request = ShiftWorkerAssignmentRequest(
            shift_worker_type="telecontrollo", date_from=record.work_date, date_to=record.work_date
        )
        record_shift_assignment(db, record, request, origin=origin)
        punches = PresenzeDailyPunch(
            daily_record_id=rid, sequence=1, entry_time=time(14), exit_time=time(21)
        )
        db.add(punches)
        db.flush()
        row = ExportTimesheetRow(db.get(PresenzeCollaborator, cid), [record], {})
        workbook = Workbook()
        detail, summary = workbook.active, workbook.create_sheet()
        record.meal_voucher_manual = True
        write_archive2_daily_values(detail, 5, row)
        write_archivio_summary_values(summary, 2, row, period_start=date(2026, 10, 1))
        assert detail.cell(5, 349).value == summary.cell(2, 23).value == 1
        record.meal_voucher_manual = False
        punches.exit_time = None
        db.flush()
        write_archive2_daily_values(detail, 5, row)
        write_archivio_summary_values(summary, 2, row, period_start=date(2026, 10, 1))
        assert detail.cell(5, 349).value == summary.cell(2, 23).value == 0


def test_non_shift_quality_accepts_imported_classification_without_provenance():
    from app.modules.presenze.models import PresenzeCollaborator
    from app.modules.presenze.services.operational_quality import build_daily_operational_quality

    user = _create_user("shift-legacy-classification")
    rid, cid = _record(user.id)
    with TestingSessionLocal() as db:
        person = db.get(PresenzeCollaborator, cid)
        person.contract_kind = "impiegato"
        record = db.get(PresenzeDailyRecord, rid)
        record.teo_minutes = 420
        classification = SimpleNamespace(ordinary_minutes=420, extra_minutes=0, source=None)
        quality = build_daily_operational_quality(
            person, record, [punch("14:00", "21:00")], classification=classification
        )
        assert quality.worked_minutes == 420
        assert quality.status == "unknown"
        assert quality.formula_code is None
        assert quality.expected_minutes == 420
        assert not any("Classificazione GAIA:" in note for note in quality.notes)


@pytest.mark.parametrize("role", ["viewer", "operator"])
def test_owner_cannot_assign_employee_wide_shift_range(role):
    owner = _create_user("shift-range-owner", role=role)
    rid, _ = _record(owner.id)
    response = client.post(
        f"/presenze/giornaliere/{rid}/turnista",
        headers={"Authorization": "Bearer " + _login(owner.username)},
        json={"shift_worker_type": "acquaiolo", "date_from": "2026-10-01", "date_to": "2026-10-31"},
    )
    assert response.status_code == 403
    with TestingSessionLocal() as db:
        assert db.scalars(select(PresenzeShiftAssignment)).all() == []
