from __future__ import annotations

import os
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from threading import Barrier
from types import SimpleNamespace
from uuid import uuid4

import pytest
from alembic.autogenerate import compare_metadata
from alembic.migration import MigrationContext
from alembic.operations import Operations
from alembic.util import load_python_file
from openpyxl import Workbook
from sqlalchemy import Column, Table, Uuid, create_engine, event, func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
from sqlalchemy.schema import CreateSchema, DropSchema

from app.modules.ruolo.models import RuoloTributiRegisteredMail
from app.modules.ruolo.notice_import_models import NoticeImportBatch, NoticeImportRow
from app.modules.ruolo.notice_import_schemas import ImportSnapshot
from app.modules.ruolo.notice_register_models import (
    NoticeAttempt,
    NoticeAudit,
    NoticeDocument,
    NoticeEvidence,
    NoticeNotification,
    NoticePosition,
    NoticeRecovery,
)
from app.modules.ruolo.services import notice_import as importer
from app.modules.ruolo.services import notice_import_excel as parser
from app.modules.ruolo.services import notice_import_poste as poste
from app.modules.ruolo.services import notice_import_register as publisher

from .test_notice_register_api import _api_metadata, _headers, _seed_auth
from .test_notice_register_api import api as api

URL = "/ruolo/tributi/registro-avvisi/importazioni"


def _metadata():
    metadata = _api_metadata()
    metadata.remove(metadata.tables["ruolo_tributi_registered_mails"])
    for name in ("ruolo_tributi_posta_online_import_jobs", "ruolo_tributi_payments"):
        Table(name, metadata, Column("id", Uuid, primary_key=True))
    for model in (RuoloTributiRegisteredMail, NoticeImportBatch, NoticeImportRow):
        model.__table__.to_metadata(metadata)
    return metadata


@pytest.fixture(scope="module", params=["sqlite", "postgresql"])
def api_engine(request):
    if request.param == "sqlite":
        engine = create_engine(
            "sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False}
        )

        @event.listens_for(engine, "connect")
        def foreign_keys(connection, _):
            connection.execute("PRAGMA foreign_keys=ON")

        _metadata().create_all(engine)
        yield engine
        engine.dispose()
        return
    url = os.getenv("GAIA_TEST_POSTGRES_URL")
    if not url:
        pytest.skip("GAIA_TEST_POSTGRES_URL non configurato")
    admin = create_engine(url)
    schema = f"notice_import_{uuid4().hex}"
    with admin.begin() as connection:
        connection.execute(CreateSchema(schema))
    engine = create_engine(url, connect_args={"options": f"-csearch_path={schema}"})
    try:
        _metadata().create_all(engine)
        yield engine
    finally:
        engine.dispose()
        with admin.begin() as connection:
            connection.execute(DropSchema(schema, cascade=True))
        admin.dispose()


def values(**columns):
    row = [None] * 63
    for index, value in {
        2: "020220001834880",
        3: "020230011722820",
        20: "TESTCF",
        21: "12024222303799",
        56: 2023,
    }.items():
        row[index] = value
    for column, value in columns.items():
        from openpyxl.utils.cell import column_index_from_string

        row[column_index_from_string(column) - 1] = value
    return row


def workbook(rows=None, sheet="Dati", headers=None):
    book = Workbook()
    page = book.active
    page.title = sheet
    header = [None] * 63
    for index, label in parser.HEADERS.items():
        header[index] = label
    page.append(header if headers is None else headers)
    for row in rows if rows is not None else [values()]:
        page.append(row)
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def preview(api, content=None):
    response = api.client.post(
        f"{URL}/excel",
        files={"file": ("avvisi.xlsx", content if content is not None else workbook())},
    )
    assert response.status_code == 201, response.text
    return response.json()


def confirm(api, batch, **extra):
    return api.client.post(
        f"{URL}/{batch['id']}/conferma",
        json={
            "digest": batch["digest"],
            "confirmed": True,
            "reason": "Verifica import storico",
            **extra,
        },
    )


def rows(api, batch):
    response = api.client.get(f"{URL}/{batch['id']}/righe")
    assert response.status_code == 200, response.text
    return response.json()["items"]


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (None, None),
        ("---", None),
        (" - ", None),
        (" ", None),
        (123.0, "123"),
        (12.5, "12.5"),
        ("0012", "0012"),
    ],
)
def test_references(value, expected):
    assert parser.reference(value) == expected


@pytest.mark.parametrize(
    ("value", "normalized", "warning"),
    [
        (None, None, None),
        ("29/06/204", "2024-06-29", "data_corretta_29_06_204"),
        (datetime(2024, 6, 29), "2024-06-29", None),
        ("30/02/2024", None, "data_non_valida"),
        ("NOTIFICATO", None, "data_non_valida"),
        (date.today() + timedelta(days=1), None, "data_non_valida"),
        ("1899-01-01", None, "data_non_valida"),
        ("2024-01-02", "2024-01-02", None),
    ],
)
def test_dates(value, normalized, warning):
    assert parser.normalized_date(value) == (normalized, warning)


def test_excel_conservative_normalization():
    content = workbook(
        [
            values(I="29/06/204", J=date(2024, 7, 1), H="NON NOTIFICATO", AX="SI"),
            values(V="different", C=202200001.0, D="---", U=None),
        ]
    )
    parsed, summary = parser.parse_excel(content)
    first = parsed[0]
    assert first["payload"]["original"]["I"] == "29/06/204"
    assert first["payload"]["dates"]["I"] == "2024-06-29"
    assert "date_discordanti" in first["anomalies"]
    assert first["payload"]["motivation"] == "NON NOTIFICATO"
    assert "notifica_da_verificare" in first["anomalies"]
    assert "riferimento_2022_numerico_verificare_zeri" in parsed[1]["anomalies"]
    assert summary["rows"] == 2
    assert summary["annual_references"] == 3
    broken = parser.parse_row(values(C="x" * 101, D=None, U="x" * 21, V=None, I="not a date"), 3)
    assert broken["payload"]["positions"] == []
    assert broken["payload"]["document_number"].startswith("Senza numero")
    assert "posizioni_assenti" in broken["anomalies"]
    assert parser.cell_value(date(2024, 1, 1)) == "2024-01-01"


@pytest.mark.parametrize(
    "content",
    [
        b"",
        b"not zip",
        workbook(sheet="wrong"),
        workbook(headers=["wrong"]),
        workbook(rows=[]),
        workbook(rows=[values(BE=2024)]),
    ],
)
def test_rejects_invalid_workbooks(content):
    with pytest.raises(ValueError):
        parser.parse_excel(content)


def test_limits(monkeypatch):
    monkeypatch.setattr(parser, "MAX_BYTES", 1)
    with pytest.raises(ValueError, match="12 MiB"):
        parser.parse_excel(b"xx")
    monkeypatch.setattr(parser, "MAX_BYTES", 1000000)
    monkeypatch.setattr(parser, "MAX_ROWS", 1)
    with pytest.raises(ValueError, match="limiti"):
        parser.parse_excel(workbook())

    class BigZip:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            pass

        def infolist(self):
            return [SimpleNamespace(file_size=65 * 1024 * 1024)]

    monkeypatch.setattr(parser, "ZipFile", lambda _: BigZip())
    with pytest.raises(ValueError, match="64 MiB"):
        parser.parse_excel(b"zip")


def test_malformed_xml_and_missing_dimensions(monkeypatch):
    with pytest.raises(ValueError, match="Dimensioni"):
        parser._read_sheet(SimpleNamespace(max_row=None, max_column=63))
    with pytest.raises(ValueError, match="Dimensioni"):
        parser._read_sheet(SimpleNamespace(max_row=3, max_column=None))
    content = workbook()

    def malformed(*args, **kwargs):
        raise SyntaxError("invalid XML")

    with monkeypatch.context() as patch:
        patch.setattr(parser, "load_workbook", malformed)
        with pytest.raises(ValueError, match="XLSX"):
            parser.parse_excel(content)
    monkeypatch.setattr(parser, "_read_sheet", malformed)
    with pytest.raises(ValueError, match="XML"):
        parser.parse_excel(content)


def test_preview_confirm_and_duplicate_content(api):
    content = workbook([values(I="29/06/204"), values(V="second")])
    batch = preview(api, content)
    assert batch["status"] == "preview"
    assert batch["summary"]["outcomes"] == {"new": 2}
    with api.session() as db:
        assert db.scalar(select(func.count()).select_from(NoticeDocument)) == 0
    assert api.client.get(f"{URL}/{batch['id']}/originale").content == content
    assert preview(api, content)["id"] == batch["id"]
    assert rows(api, batch)[0]["payload"]["original"]["I"] == "29/06/204"
    assert confirm(api, batch, confirmed=False).status_code == 422
    assert confirm(api, batch, actor_id=2).status_code == 422
    assert confirm(api, batch, reason="   ").status_code == 422
    assert confirm(api, batch, digest="wrong").status_code == 409
    result = confirm(api, batch)
    assert result.status_code == 200, result.text
    assert result.json()["summary"]["outcomes"] == {"imported": 2}
    repeated = confirm(api, batch).json()
    assert repeated["summary"] == result.json()["summary"]
    assert repeated["confirmed_at"].rstrip("Z") == result.json()["confirmed_at"].rstrip("Z")
    with api.session() as db:
        assert db.scalar(select(func.count()).select_from(NoticeDocument)) == 2
        assert db.scalar(select(func.count()).select_from(NoticePosition)) == 4
        assert set(db.scalars(select(NoticeNotification.state))) == {"da_verificare"}
        assert set(db.scalars(select(NoticeRecovery.state))) == {"da_verificare"}
        assert set(db.scalars(select(NoticePosition.avviso_id))) == {None}
        assert db.scalar(select(func.count()).select_from(NoticeEvidence)) == 2
        assert set(db.scalars(select(NoticeAudit.actor_id))) == {1}
    assert api.client.get(URL).json()["total"] == 1
    assert api.client.get(f"{URL}/{batch['id']}").json()["status"] == "confirmed"
    assert (
        api.client.get(f"{URL}/{batch['id']}/righe?page=2&page_size=1").json()["items"][0][
            "row_number"
        ]
        == 3
    )
    # A different binary workbook with identical row values must not duplicate documents.
    other = preview(api, workbook([values(I="29/06/204"), values(V="second"), [None] * 63]))
    assert other["summary"]["outcomes"] == {"duplicate": 2}
    assert confirm(api, other).json()["summary"]["outcomes"] == {"duplicate": 2}


def test_conflicts_and_repeated_rows_do_not_overwrite(api):
    batch = preview(
        api, workbook([values(), values(), values(H="different"), values(H="different")])
    )
    assert batch["summary"]["outcomes"] == {"new": 1, "duplicate": 1, "conflict": 2}
    result = confirm(api, batch)
    assert result.status_code == 200, result.text
    assert result.json()["summary"]["outcomes"] == {"imported": 1, "duplicate": 1, "conflict": 2}
    with api.session() as db:
        document = db.scalar(select(NoticeDocument))
        document.document_number = "Operatore corretto"
        db.commit()
    changed = preview(api, workbook([values(H="different"), values(H="different")]))
    assert changed["summary"]["outcomes"] == {"conflict": 2}
    assert confirm(api, changed).json()["summary"]["outcomes"] == {"conflict": 2}
    assert rows(api, changed)[0]["document_id"] is not None
    with api.session() as db:
        assert db.scalar(select(NoticeDocument.document_number)) == "Operatore corretto"


def test_confirmation_rollback(api, monkeypatch):
    batch = preview(api)
    original = publisher._evidence

    def fail(*args):
        original(*args)
        raise ValueError("fallimento simulato")

    monkeypatch.setattr(publisher, "_evidence", fail)
    assert confirm(api, batch).status_code == 422
    with api.session() as db:
        assert db.scalar(select(func.count()).select_from(NoticeDocument)) == 0
        assert db.scalar(select(NoticeImportBatch.status)) == "preview"


@pytest.mark.parametrize("user_id", [2, 3, 4, 5])
def test_import_authorization(api, user_id):
    batch = preview(api)
    api.client.headers.update(_headers(user_id))
    denied = 401 if user_id == 4 else 403
    for path in ("/excel", "/poste", f"/{batch['id']}/conferma"):
        assert (
            api.client.post(URL + path, files={"file": ("a.xlsx", workbook())}).status_code
            == denied
        )
    if user_id != 2:
        for path in ("", f"/{batch['id']}", f"/{batch['id']}/righe", f"/{batch['id']}/originale"):
            assert api.client.get(URL + path).status_code == denied
    else:
        assert api.client.get(URL).status_code == 200


def test_import_not_found_and_invalid_input(api):
    unknown = {"id": str(uuid4()), "digest": "x"}
    assert confirm(api, unknown).status_code == 404
    for suffix in ("", "/righe", "/originale"):
        assert api.client.get(f"{URL}/{unknown['id']}{suffix}").status_code == 404
    assert api.client.post(f"{URL}/excel", files={"file": ("bad", b"bad")}).status_code == 422
    assert api.client.post(f"{URL}/poste").status_code == 422
    assert api.client.get(f"{URL}?page=0").status_code == 422


def test_poste_preserves_source_without_inferring_notification(api, monkeypatch):
    with api.session() as db:
        mails = [
            RuoloTributiRegisteredMail(
                id=uuid4(),
                source_shipment_id=f"post-{i}",
                recipient_index=0,
                tracking_number=f"tracking-{i}",
                status_label="Servizio erogato",
                annualita_json=[2022, 2023],
                sent_at=datetime(2024, 1, 1) if i == 0 else None,
                raw_payload_json={"info": "originale"},
            )
            for i in range(2)
        ]
        db.add_all(mails)
        db.commit()
    response = api.client.post(f"{URL}/poste")
    assert response.status_code == 201, response.text
    batch = response.json()
    assert confirm(api, batch).status_code == 200
    with api.session() as db:
        assert db.scalar(select(func.count()).select_from(NoticeAttempt)) == 2
        assert db.scalar(select(func.count()).select_from(NoticePosition)) == 0
        assert set(db.scalars(select(NoticeNotification.state))) == {"da_verificare"}
        assert (
            db.scalar(select(NoticeEvidence)).original_json["original"]["status_label"]
            == "Servizio erogato"
        )
        monkeypatch.setattr(poste, "MAX_ROWS", 1)
        with pytest.raises(ValueError, match="10000"):
            poste.snapshot_poste(db)
        monkeypatch.setattr(poste, "MAX_ROWS", 10000)
        monkeypatch.setattr(
            poste,
            "json_bytes",
            lambda obj: (
                b" " * (65 * 1024 * 1024) if isinstance(obj, list) else parser.json_bytes(obj)
            ),
        )
        with pytest.raises(ValueError, match="64 MiB"):
            poste.snapshot_poste(db)


def test_import_migration_roundtrip(api):
    with api.session() as db:
        connection = db.connection()
        metadata = _metadata()
        # This test characterizes the import revision before conflict decisions existed.
        table = metadata.tables[NoticeImportRow.__tablename__]
        table._columns.remove(table.c.resolution)
        tables = [
            metadata.tables[model.__tablename__] for model in (NoticeImportRow, NoticeImportBatch)
        ]
        metadata.drop_all(connection, tables=tables)
        path = Path(__file__).resolve().parents[2] / "alembic" / "versions"
        revision = load_python_file(str(path), "20260918_0900_ruolo_notice_import.py")
        assert revision.down_revision == "20260917_0900"
        with Operations.context(MigrationContext.configure(connection)):
            revision.upgrade()
            assert compare_metadata(MigrationContext.configure(connection), metadata) == []
            revision.downgrade()
            revision.upgrade()
            assert compare_metadata(MigrationContext.configure(connection), metadata) == []


def test_integrity_error_at_commit_rolls_back(api, monkeypatch):
    batch = preview(api)

    def failed_commit(self):
        raise IntegrityError("commit", {}, Exception("private database error"))

    with monkeypatch.context() as patch:
        patch.setattr(Session, "commit", failed_commit)
        response = confirm(api, batch)
    assert response.status_code == 409
    assert "private" not in response.text
    with api.session() as db:
        assert db.scalar(select(func.count()).select_from(NoticeDocument)) == 0
        assert db.scalar(select(NoticeImportBatch.status)) == "preview"


@pytest.mark.parametrize("separate_batches", [False, True])
def test_concurrent_confirmations_do_not_duplicate(api_engine, separate_batches):
    if api_engine.dialect.name != "postgresql":
        pytest.skip("Concorrenza verificata su PostgreSQL, non simulata da SQLite")
    schema = f"notice_race_{uuid4().hex}"
    with api_engine.begin() as connection:
        connection.execute(CreateSchema(schema))
    engine = create_engine(api_engine.url, connect_args={"options": f"-csearch_path={schema}"})
    try:
        with engine.begin() as connection:
            _metadata().create_all(connection)
            _seed_auth(connection)
        content = workbook()
        parsed, summary = parser.parse_excel(content)
        with Session(engine) as db:
            first = importer.stage_import(
                db, ImportSnapshot("excel_2022_2023", "one.xlsx", content, parsed, summary), 1
            )
            second = first
            if separate_batches:
                second = importer.stage_import(
                    db,
                    ImportSnapshot(
                        "excel_2022_2023", "two.xlsx", content + b"variant", parsed, summary
                    ),
                    1,
                )
            keys = [(item.id, item.digest) for item in (first, second)]
            db.commit()
        barrier = Barrier(2)

        def publish(key):
            barrier.wait(timeout=10)
            with Session(engine) as db:
                try:
                    importer.confirm_import(db, key[0], key[1], 1, "Conferma concorrente")
                    db.commit()
                except IntegrityError:
                    db.rollback()
                    importer.confirm_import(db, key[0], key[1], 1, "Retry concorrente")
                    db.commit()

        with ThreadPoolExecutor(max_workers=2) as pool:
            list(pool.map(publish, keys))
        with Session(engine) as db:
            assert db.scalar(select(func.count()).select_from(NoticeDocument)) == 1
            assert db.scalar(select(func.count()).select_from(NoticeAudit)) == 1
            assert set(db.scalars(select(NoticeImportBatch.status))) == {"confirmed"}
    finally:
        engine.dispose()
        with api_engine.begin() as connection:
            connection.execute(DropSchema(schema, cascade=True))
