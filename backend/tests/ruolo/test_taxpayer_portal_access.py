"""Real workbook validation, transactional import and encrypted printed snapshots."""

import json
import runpy
import sys
from pathlib import Path

import pytest
from alembic.migration import MigrationContext
from alembic.operations import Operations
from alembic.util import load_python_file
from cryptography.fernet import InvalidToken
from openpyxl import Workbook
from sqlalchemy import create_engine, inspect, select
from sqlalchemy.orm import Session

from app.core import database
from app.modules.ruolo.portal_access_models import RuoloTaxpayerPortalAccess as Access
from app.modules.ruolo.services.notice_branding import notice_portal_html
from app.modules.ruolo.services.notice_candidates import payload_overrides
from app.modules.ruolo.services.taxpayer_portal_access import (
    PORTAL_ACCESS_KEY,
    printable_portal_access,
    with_portal_access,
)
from app.modules.ruolo.services.taxpayer_portal_import import (
    PortalAccessRow,
    import_access_rows,
    read_access_workbook,
)
from app.scripts import import_taxpayer_portal_access as script
from app.services.elaborazioni_credentials import get_credential_fernet

CF = "RSSMRA80A01H501Z"
OTHER_CF = "VRDLGU80A01H501W"


@pytest.fixture
def portal_engine():
    engine = create_engine("sqlite://")
    Access.__table__.create(engine)
    yield engine
    engine.dispose()


def _workbook(tmp_path, rows, *, headers=("CodFisc", "Utenza", "PIN")):
    path = tmp_path / "access.xlsx"
    workbook = Workbook()
    workbook.active.append(headers)
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)
    workbook.close()
    return path


def _snapshot(*pairs):
    fernet = get_credential_fernet()
    return {
        PORTAL_ACCESS_KEY: [
            {"account_code": code, "pin_encrypted": fernet.encrypt(pin.encode()).decode()}
            for code, pin in pairs
        ]
    }


def test_excel_keeps_string_and_formatted_numeric_leading_zeros_and_duplicate_rows(tmp_path):
    path = _workbook(tmp_path, [(f" {CF.lower()} ", " 0001234 ", " 0012345 "), (CF, 123, 456)])
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    workbook.active["B3"].number_format = "0000000"
    workbook.active["C3"].number_format = "0000000"
    workbook.active.append((CF, "0001234", "0012345"))
    workbook.active.append((None, None, None, "Totale"))
    workbook.save(path)
    workbook.close()
    rows, report = read_access_workbook(path)
    assert rows == [
        PortalAccessRow(CF, "0001234", "0012345"),
        PortalAccessRow(CF, "0000123", "0000456"),
    ]
    assert report == {
        "source_rows": 3,
        "duplicate_rows": 1,
        "summary_or_blank_rows": 1,
        "accounts": 2,
        "tax_codes": 1,
    }
    assert "0012345" not in repr(rows)


@pytest.mark.parametrize(
    "rows,headers",
    [
        ([], ("CodFisc", "Utenza", "Utenza", "PIN")),
        ([], ("CodFisc", "Utenza")),
        ([], ()),
        ([], ("CodFisc", "Utenza", "PIN")),
        ([("bad", "123", "7654321")], ("CodFisc", "Utenza", "PIN")),
        ([("\N{FULLWIDTH DIGIT ZERO}" * 11, "123", "7654321")], ("CodFisc", "Utenza", "PIN")),
        ([("A" * 10 + "!", "123", "7654321")], ("CodFisc", "Utenza", "PIN")),
        ([(None, "123", "7654321")], ("CodFisc", "Utenza", "PIN")),
        ([(None, None, 1234567)], ("CodFisc", "Utenza", "PIN")),
    ],
)
def test_excel_rejects_missing_headers_empty_sources_and_invalid_tax_codes(tmp_path, rows, headers):
    path = _workbook(tmp_path, rows, headers=headers)
    with pytest.raises(ValueError):
        read_access_workbook(path)


@pytest.mark.parametrize(
    "value",
    [
        None,
        True,
        1.5,
        -1,
        "\N{FULLWIDTH DIGIT ONE}\N{FULLWIDTH DIGIT TWO}\N{FULLWIDTH DIGIT THREE}",
        " ",
        "ab",
        "1" * 33,
        "=1+1",
    ],
)
def test_invalid_pin_is_rejected_without_echoing_its_value(tmp_path, value):
    path = _workbook(tmp_path, [(CF, "1234567", value)])
    with pytest.raises(ValueError, match="Credenziali non valide alla riga 2"):
        read_access_workbook(path)


@pytest.mark.parametrize("second", [(CF, "1234567", "2222222"), (OTHER_CF, "1234567", "1111111")])
def test_excel_refuses_conflicting_pin_or_tax_code_for_the_same_account(tmp_path, second):
    path = _workbook(tmp_path, [(CF, "1234567", "1111111"), second])
    with pytest.raises(ValueError, match="Codice ditta in conflitto alla riga 3"):
        read_access_workbook(path)


def test_import_is_encrypted_idempotent_updates_pin_and_rolls_back(portal_engine):
    first = PortalAccessRow(CF, "0000123", "0012345")
    second = PortalAccessRow(CF, "0000456", "0076543")
    with Session(portal_engine) as db:
        assert import_access_rows(db, [first, second], "a" * 64) == {
            "created": 2,
            "updated": 0,
            "unchanged": 0,
        }
        db.commit()
        before = db.get(Access, first.account_code).pin_encrypted
        assert before != first.pin.encode()
        assert get_credential_fernet().decrypt(before).decode() == first.pin
        assert import_access_rows(db, [first, second], "a" * 64) == {
            "created": 0,
            "updated": 0,
            "unchanged": 2,
        }
        assert db.get(Access, first.account_code).pin_encrypted == before
        assert import_access_rows(
            db, [PortalAccessRow(CF, first.account_code, "1234567")], "b" * 64
        ) == {"created": 0, "updated": 1, "unchanged": 0}
        assert db.get(Access, first.account_code).source_sha256 == "b" * 64
        db.rollback()
        assert db.get(Access, first.account_code).pin_encrypted == before
        assert import_access_rows(db, [], "a" * 64) == {"created": 0, "updated": 0, "unchanged": 0}


def test_import_checks_existing_account_ownership_before_any_write(portal_engine):
    with Session(portal_engine) as db:
        import_access_rows(db, [PortalAccessRow(CF, "1234567", "1111111")], "a" * 64)
        db.commit()
        with pytest.raises(ValueError, match="già associato"):
            import_access_rows(
                db,
                [
                    PortalAccessRow(CF, "7654321", "2222222"),
                    PortalAccessRow(OTHER_CF, "1234567", "1111111"),
                ],
                "b" * 64,
            )
        assert len(list(db.scalars(select(Access)))) == 1


def test_snapshot_prints_only_an_unambiguous_account_of_the_exact_recipient(portal_engine):
    with Session(portal_engine) as db:
        import_access_rows(
            db,
            [
                PortalAccessRow(CF, "2222222", "0022222"),
                PortalAccessRow(CF, "1111111", "0011111"),
                PortalAccessRow("01234567890", "3333333", "0033333"),
            ],
            "a" * 64,
        )
        candidates = [
            {"codice_fiscale": f" {CF.lower()} ", "display_name": "Rossi"},
            {"codice_fiscale": OTHER_CF, "display_name": "Rossi"},
            {"codice_fiscale": "01234567890"},
        ]
        enriched = with_portal_access(db, candidates)
        payload = payload_overrides(enriched[2])
        assert "0011111" not in json.dumps(payload)
        assert enriched[0][PORTAL_ACCESS_KEY] == []
        assert printable_portal_access(enriched[0]) == []
        assert printable_portal_access(enriched[1]) == []
        assert printable_portal_access(payload) == [("3333333", "0033333")]
        assert "0033333" not in json.dumps(payload)
        assert PORTAL_ACCESS_KEY not in candidates[0]
        assert with_portal_access(db, []) == []
        assert (
            printable_portal_access({}) == printable_portal_access({PORTAL_ACCESS_KEY: None}) == []
        )
        import_access_rows(db, [PortalAccessRow("01234567890", "3333333", "9999999")], "b" * 64)
        assert printable_portal_access(payload) == [("3333333", "0033333")]


def test_printed_portal_has_one_credential_or_office_fallback_and_escapes_html():
    rendered = notice_portal_html(_snapshot(("0000123", "0012345")))
    assert rendered.count('class="access-row"') == 1
    assert "0000123" in rendered and "0012345" in rendered
    assert "Cartella contribuenza" in rendered and "contattare gli uffici" not in rendered
    assert "contattare gli uffici del Consorzio" in notice_portal_html({})
    escaped = notice_portal_html(_snapshot(("<account>", "<pin>")))
    assert "&lt;account&gt;" in escaped and "&lt;pin&gt;" in escaped
    assert "<pin>" not in escaped


def test_historical_multi_account_snapshot_uses_office_fallback_without_decrypting(monkeypatch):
    payload = _snapshot(("0000123", "0012345"), ("0000456", "0076543"))

    def unexpected_decryption():
        pytest.fail("An ambiguous historical snapshot must not decrypt any credentials")

    monkeypatch.setattr(
        "app.modules.ruolo.services.taxpayer_portal_access.get_credential_fernet",
        unexpected_decryption,
    )
    rendered = notice_portal_html(payload)
    assert 'class="access-row"' not in rendered
    assert "contattare gli uffici del Consorzio" in rendered
    assert all(value not in rendered for value in ("0000123", "0012345", "0000456", "0076543"))


def test_corrupt_snapshot_does_not_print_an_unverified_pin():
    with pytest.raises(InvalidToken):
        printable_portal_access(
            {PORTAL_ACCESS_KEY: [{"account_code": "123", "pin_encrypted": "corrupt"}]}
        )


def test_cli_dry_run_apply_repeat_and_module_entrypoint(
    tmp_path, portal_engine, monkeypatch, capsys
):
    path = _workbook(tmp_path, [(CF, "0000123", 1234567)])

    def factory():
        return Session(portal_engine)

    monkeypatch.setattr(script, "SessionLocal", factory)
    script.main([str(path)])
    report = json.loads(capsys.readouterr().out)
    assert report["created"] == 1 and report["applied"] is False
    with Session(portal_engine) as db:
        assert list(db.scalars(select(Access))) == []
    script.main([str(path), "--apply"])
    assert json.loads(capsys.readouterr().out)["applied"] is True
    monkeypatch.setattr(database, "SessionLocal", factory)
    monkeypatch.setattr(sys, "argv", ["import_taxpayer_portal_access", str(path), "--apply"])
    monkeypatch.delitem(sys.modules, "app.scripts.import_taxpayer_portal_access")
    runpy.run_module("app.scripts.import_taxpayer_portal_access", run_name="__main__")
    report = json.loads(capsys.readouterr().out)
    assert report["unchanged"] == 1 and "1234567" not in json.dumps(report)


def test_portal_migration_matches_model_and_roundtrips():
    revision = load_python_file(
        str(Path(__file__).parents[2] / "alembic" / "versions"),
        "20261027_1100_taxpayer_portal_access.py",
    )
    engine = create_engine("sqlite://")
    with engine.begin() as connection, Operations.context(MigrationContext.configure(connection)):
        revision.upgrade()
        columns = {c["name"] for c in inspect(connection).get_columns(Access.__tablename__)}
        assert columns == set(Access.__table__.columns.keys())
        assert inspect(connection).get_pk_constraint(Access.__tablename__)[
            "constrained_columns"
        ] == ["account_code"]
        assert inspect(connection).get_indexes(Access.__tablename__)[0]["column_names"] == [
            "tax_code"
        ]
        revision.downgrade()
        assert Access.__tablename__ not in inspect(connection).get_table_names()
    engine.dispose()
