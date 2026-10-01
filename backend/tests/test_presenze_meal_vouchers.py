from datetime import date
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from test_presenze_api import (  # noqa: F401 - shared API fixture
    TestingSessionLocal,
    _create_user,
    _login,
    client,
    setup_database,
)

from app.modules.presenze.models import PresenzeCollaborator, PresenzeDailyRecord
from app.modules.presenze.services import import_jobs
from app.modules.presenze.services.meal_vouchers import meal_voucher_values
from app.modules.presenze.services.parser import ParsedCollaboratorPayload, ParsedImportPayload
from app.modules.presenze.services.xlsm_export import (
    ExportTimesheetRow,
    write_archive2_daily_values,
    write_archivio_summary_values,
)


@pytest.mark.parametrize(
    "day,extra,manual,count,sources",
    [
        (date(2026, 8, 25), 240, False, 0, []),
        (date(2026, 8, 26), 119, False, 0, []),
        (date(2026, 8, 26), 120, False, 1, ["automatic"]),
        (date(2026, 8, 26), 0, True, 1, ["manual"]),
        (date(2026, 8, 26), 120, True, 1, ["automatic", "manual"]),
        (date(2026, 8, 25), None, True, 1, ["manual"]),
    ],
)
def test_daily_count_and_provenance(day, extra, manual, count, sources):
    values = meal_voucher_values(SimpleNamespace(work_date=day, meal_voucher_manual=manual), extra)
    assert values["meal_voucher_count"] == count
    assert values["meal_voucher_sources"] == sources


def _record(owner_id, *, extra=0):
    with TestingSessionLocal() as db:
        person = PresenzeCollaborator(
            owner_user_id=owner_id,
            employee_code="1407",
            name="Persona test",
            contract_kind="operaio",
            operai_group="agrario",
        )
        db.add(person)
        db.flush()
        day = PresenzeDailyRecord(
            collaborator_id=person.id,
            owner_user_id=owner_id,
            work_date=date(2026, 10, 1),
            schedule_code="OPE0613",
            ordinary_minutes=420,
            mpe_minutes=extra,
            reperibilita_unit="none",
            validation_status="pending",
        )
        db.add(day)
        db.commit()
        db.refresh(day)
        return day.id, person.id


def test_api_audits_idempotent_grant_and_revoke_preserving_automatic():
    user = _create_user("meal-admin")
    rid, _ = _record(user.id, extra=120)
    headers = {"Authorization": "Bearer " + _login(user.username)}
    for enabled in [True, True, False]:
        response = client.patch(
            "/presenze/giornaliere/" + str(rid),
            headers=headers,
            json={"meal_voucher_manual": enabled},
        )
        assert response.status_code == 200, response.text
        assert response.json()["meal_voucher_count"] == 1
    result = response.json()
    assert result["meal_voucher_sources"] == ["automatic"]
    assert [e["enabled"] for e in result["meal_voucher_audit"]] == [True, False]
    assert all(
        e["actor_user_id"] == user.id and e["source"] == "gaia_web"
        for e in result["meal_voucher_audit"]
    )
    matrix = client.get(
        "/presenze/giornaliere/matrix",
        headers=headers,
        params={"date_from": "2026-10-01", "date_to": "2026-10-31"},
    )
    assert matrix.status_code == 200
    assert matrix.json()["items"][0]["meal_voucher_count"] == 1
    # Full list and matrix expose the same daily benefit.
    listing = client.get(
        "/presenze/giornaliere",
        headers=headers,
        params={"date_from": "2026-10-01", "date_to": "2026-10-31"},
    )
    assert listing.status_code == 200
    assert listing.json()["items"][0]["meal_voucher_count"] == 1


def test_api_manual_permission_and_invalid_null():
    owner = _create_user("meal-owner")
    rid, _ = _record(owner.id)
    reader = _create_user("meal-reader", role="viewer")
    headers = {"Authorization": "Bearer " + _login(reader.username)}
    response = client.patch(
        "/presenze/giornaliere/" + str(rid), headers=headers, json={"meal_voucher_manual": True}
    )
    assert response.status_code in {403, 404}
    headers = {"Authorization": "Bearer " + _login(owner.username)}
    assert (
        client.patch(
            "/presenze/giornaliere/" + str(rid), headers=headers, json={"meal_voucher_manual": None}
        ).status_code
        == 422
    )


def test_reimport_preserves_manual_audit_replaces_punches_and_export_counts():
    user = _create_user("meal-import")
    rid, cid = _record(user.id)
    headers = {"Authorization": "Bearer " + _login(user.username)}
    response = client.patch(
        "/presenze/giornaliere/" + str(rid), headers=headers, json={"meal_voucher_manual": True}
    )
    assert response.status_code == 200
    audit = response.json()["meal_voucher_audit"]
    payload = ParsedCollaboratorPayload(
        collaborator={"employee_code": "1407", "name": "Persona test"},
        company_label=None,
        period_start=date(2026, 10, 1),
        period_end=date(2026, 10, 31),
        daily_rows=[
            {
                "work_date": "01/10/2026",
                "ordinary": "07:00",
                "punches": [{"entry": "06:00", "exit": "14:00"}],
            }
        ],
        summary_rows=[],
    )
    with TestingSessionLocal() as db:
        parsed = ParsedImportPayload(
            period_start=payload.period_start,
            period_end=payload.period_end,
            collaborators=[payload],
            errors=[],
        )
        job = import_jobs.create_import_job(
            db, parsed=parsed, requested_by_user_id=user.id, filename="corrected.json"
        )
        import_jobs.import_collaborator_payload(db, payload=payload, job=job)
        import_jobs.finalize_import_job(db, job=job)
        db.commit()
        record = db.get(PresenzeDailyRecord, rid)
        assert record.meal_voucher_manual and record.meal_voucher_audit == audit
        assert len(record.raw_payload_json["punches"]) == 1
        person = db.get(PresenzeCollaborator, cid)
        row = ExportTimesheetRow(person, [record], {})
        wb = Workbook()
        detail = wb.active
        summary = wb.create_sheet()
        write_archive2_daily_values(detail, 5, row)
        write_archivio_summary_values(summary, 2, row, period_start=date(2026, 10, 1))
        assert detail.cell(5, 349).value == 1
        assert summary.cell(2, 23).value == 1
        record.mpe_minutes = 120
        write_archive2_daily_values(detail, 5, row)
        write_archivio_summary_values(summary, 2, row, period_start=date(2026, 10, 1))
        assert detail.cell(5, 349).value == summary.cell(2, 23).value == 1
        record.meal_voucher_manual = False
        record.mpe_minutes = 0
        write_archive2_daily_values(detail, 5, row)
        assert detail.cell(5, 349).value == 0
