from __future__ import annotations

import asyncio
import json
from datetime import datetime
from io import BytesIO

import httpx
import pytest
from openpyxl import load_workbook

from app.modules.elaborazioni.capacitas.apps.inbollettini import exports as boll
from app.modules.elaborazioni.capacitas.export_http import (
    ExportStepError,
    SessionExpiredError,
    ensure_portal_response,
    ensure_xlsx,
    kendo_grid_rows,
    parse_select_options,
)
from app.modules.elaborazioni.capacitas.export_workbook import (
    cell_value,
    collect_columns,
    rows_to_payload,
    xlsx_data_rows,
)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def response(
    text="", *, url="https://incass3.test/pages/x.aspx", status=200, headers=None, content=None
):
    request = httpx.Request("GET", url)
    kwargs = {"content": content} if content is not None else {"text": text}
    return httpx.Response(status, request=request, headers=headers or {}, **kwargs)


def read_sheet(content: bytes) -> list[tuple]:
    workbook = load_workbook(BytesIO(content))
    return list(workbook.active.iter_rows(values_only=True))


def test_rows_to_payload_keeps_every_column_and_converts_values():
    rows = [
        {"strAvviso": "025", "dtRataTScadenza": "/Date(1793401200000)/", "dblcarico": 388.11},
        {
            "strAvviso": "026",
            "extra": {"a": 1},
            "flag": True,
            "missing": None,
            "neg": "/Date(-1000+0100)/",
        },
    ]

    payload = rows_to_payload(
        rows, sheet_title="Registro bollettini molto lungo oltre trentuno caratteri"
    )

    sheet = read_sheet(payload.content)
    assert payload.rows == 2
    assert sheet[0] == (
        "strAvviso",
        "dtRataTScadenza",
        "dblcarico",
        "extra",
        "flag",
        "missing",
        "neg",
    )
    assert sheet[1][:3] == ("025", datetime(2026, 10, 31, 0, 0), 388.11)
    assert sheet[2][3] == '{"a": 1}'
    assert sheet[2][4] is True
    assert xlsx_data_rows(payload.content) == 2


def test_collect_columns_and_cell_value_edge_cases():
    assert collect_columns([{"b": 1}, {"a": 2, "b": 3}]) == ["b", "a"]
    assert cell_value("testo") == "testo"
    assert cell_value(None) is None
    assert cell_value(["x"]) == '["x"]'


def test_xlsx_data_rows_of_empty_export_is_zero():
    assert xlsx_data_rows(rows_to_payload([], sheet_title="vuoto").content) == 0


def test_xlsx_data_rows_ignores_portal_trailing_newline_row():
    payload = rows_to_payload(
        [{"Reso": "\n   ", "Anno": "2025"}, {"Reso": "\n", "Anno": None}], sheet_title="x"
    )
    assert xlsx_data_rows(payload.content) == 1


@pytest.mark.parametrize(
    ("url", "text", "error"),
    [
        ("https://sso.servizicapacitas.com/pages/login.aspx", "", SessionExpiredError),
        ("https://incass3.test/pages/main.aspx", '<input name="txtUsername">', SessionExpiredError),
        ("https://incass3.test/pages/errore.aspx", "", ExportStepError),
        ("https://incass3.test/pages/ajax.aspx", "NOparametri errati", ExportStepError),
    ],
)
def test_ensure_portal_response_rejects_portal_failures(url, text, error):
    with pytest.raises(error):
        ensure_portal_response(response(text, url=url), step="test")


def test_ensure_portal_response_accepts_regular_page_and_raises_http_errors():
    ensure_portal_response(response("SI"), step="ok")
    with pytest.raises(httpx.HTTPStatusError):
        ensure_portal_response(response("boom", status=503), step="ko")


def test_parse_select_options_skips_blank_values():
    page = '<select name="x" id="ContentMain_ddlAnno"><option value=""></option><option value="2025">2025</option><option value="9926">9926</option></select>'
    assert parse_select_options(page, "ContentMain_ddlAnno") == ["2025", "9926"]
    with pytest.raises(ExportStepError) as excinfo:
        parse_select_options("<html></html>", "ContentMain_ddlAnno")
    assert excinfo.value.retryable is False


def test_ensure_xlsx_requires_spreadsheet_body():
    assert (
        ensure_xlsx(response(content=b"PK\x03\x04", headers={"content-type": XLSX}), step="x")
        == b"PK\x03\x04"
    )
    with pytest.raises(ExportStepError, match="text/html"):
        ensure_xlsx(response("<html>", headers={"content-type": "text/html"}), step="x")
    with pytest.raises(ExportStepError, match="senza content-type"):
        ensure_xlsx(response(content=b"PK"), step="x")


class FakeManager:
    def __init__(self, handler):
        self.http = httpx.AsyncClient(transport=httpx.MockTransport(handler))

    def get_http_client(self):
        return self.http

    def get_token(self):
        return "token-1"


def test_inbollettini_lists_esercizi_and_exports_registro():
    requests: list[httpx.Request] = []

    def handler(request):
        requests.append(request)
        if request.method == "GET":
            return httpx.Response(
                200,
                text='<select id="ContentMain_ddlEsercizio"><option value=""></option><option value="2525">2525</option></select>',
            )
        body = {
            "results": [{"strAvviso": "1", "dtDataOperazione": "/Date(0)/"}],
            "total": 1,
            "err": None,
        }
        return httpx.Response(200, text=json.dumps(body))

    client = boll.InBollettiniExportClient(FakeManager(handler))

    async def scenario():
        return await client.list_esercizi(), await client.export_registro("2525")

    esercizi, payload = asyncio.run(scenario())

    assert esercizi == ["2525"]
    assert payload.rows == 1
    assert requests[0].url.params["token"] == "token-1"
    form = dict(item.split("=", 1) for item in requests[1].content.decode().split("&"))
    assert form["op"] == "registroBoll" and form["esercizio"] == "2525" and form["cc"] == ""


@pytest.mark.parametrize(
    ("body", "message"),
    [
        ("not json", "non JSON"),
        (json.dumps([1]), "senza elenco results"),
        (json.dumps({"results": [], "err": "boom"}), "errore: boom"),
        (json.dumps({"results": [{"a": 1}], "total": 2}), "1 righe su 2"),
    ],
)
def test_kendo_grid_rows_rejects_invalid_payloads(body, message):
    with pytest.raises(ExportStepError, match=message):
        kendo_grid_rows(body, step="registro")


def test_kendo_grid_rows_ignores_non_dict_rows_without_total():
    assert kendo_grid_rows(json.dumps({"results": [{"a": 1}, "x"]}), step="r") == [{"a": 1}]


def test_inbollettini_session_expiry_is_reported():
    def handler(request):
        return httpx.Response(200, text='<input name="txtUsername">')

    client = boll.InBollettiniExportClient(FakeManager(handler))
    with pytest.raises(SessionExpiredError):
        asyncio.run(client.export_registro("2525"))
