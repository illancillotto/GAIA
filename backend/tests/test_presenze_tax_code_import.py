from datetime import date
from pathlib import Path
from types import SimpleNamespace
from uuid import uuid4

import pytest
from openpyxl import Workbook

from app.modules.presenze.models import PresenzeCollaborator, PresenzeDailyRecord
from app.modules.presenze.services.gate_mobile_payloads import _record_collaborator_values
from app.modules.presenze.services.tax_code_import import (
    TaxCodeImportEntry,
    TaxCodeImportError,
    extract_tax_codes,
    import_tax_codes,
)

CF154 = "PLLGNN63B20E972L"
CF155 = "SNNJNY73A24M168L"
CF2176 = "SCNSMN77C02G113B"


def _workbook(path: Path, *, conflict: bool = False) -> None:
    workbook = Workbook()
    operai = workbook.active
    operai.title = "Operai"
    operai.cell(2, 4).value = 154
    operai.cell(2, 16).value = CF154
    riepilogo = workbook.create_sheet("Riepilogo")
    riepilogo.cell(2, 1).value = 154
    riepilogo.cell(2, 3).value = f"2/2026-{CF154 if not conflict else CF155}"
    riepilogo.cell(3, 1).value = 155
    riepilogo.cell(3, 3).value = f"2/2026-{CF155}"
    workbook.save(path)


def test_extract_tax_codes_prefers_operai_and_falls_back_to_riepilogo(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    _workbook(path)
    entries = {item.employee_code: item for item in extract_tax_codes(path)}
    assert entries["154"].tax_code == CF154
    assert entries["154"].source == "Operai!P2"
    assert entries["155"].tax_code == CF155
    assert entries["155"].source == "Riepilogo!C3"


def test_extract_tax_codes_rejects_conflicting_sources(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    _workbook(path, conflict=True)
    with pytest.raises(ValueError, match="154"):
        extract_tax_codes(path)


class _FakeDb:
    def __init__(self, collaborators: list[object]) -> None:
        self.collaborators = collaborators
        self.committed = False
        self.rolled_back = False

    def scalars(self, _statement):
        return SimpleNamespace(all=lambda: self.collaborators)

    def commit(self) -> None:
        self.committed = True

    def rollback(self) -> None:
        self.rolled_back = True


def test_import_tax_codes_reports_and_applies_without_using_names() -> None:
    entries = [
        TaxCodeImportEntry("154", CF154, "Riepilogo!C145"),
        TaxCodeImportEntry("missing", CF155, "Riepilogo!C158"),
    ]
    collaborator = SimpleNamespace(employee_code="154", tax_code=None)
    db = _FakeDb([collaborator])
    dry_report = import_tax_codes(db, entries, apply=False)
    assert dry_report == {
        "entries": 2,
        "matched": 1,
        "missing": 1,
        "conflicts": 0,
        "changed": 1,
        "unchanged": 0,
        "applied": 0,
    }
    assert collaborator.tax_code is None
    with pytest.raises(TaxCodeImportError):
        import_tax_codes(db, entries, apply=True)
    assert collaborator.tax_code is None
    assert not db.committed
    assert db.rolled_back


def test_import_tax_codes_does_not_overwrite_existing_different_code() -> None:
    collaborator = SimpleNamespace(employee_code="154", tax_code=CF155)
    db = _FakeDb([collaborator])
    with pytest.raises(TaxCodeImportError) as error:
        import_tax_codes(db, [TaxCodeImportEntry("154", CF154, "Riepilogo!C145")], apply=True)
    assert error.value.report["conflicts"] == 1
    assert collaborator.tax_code == CF155
    assert not db.committed


def test_daily_gate_payload_carries_saved_tax_code() -> None:
    collaborator = PresenzeCollaborator(
        id=uuid4(), employee_code="154", name="Persona", tax_code=CF154
    )
    record = PresenzeDailyRecord(
        id=uuid4(), collaborator_id=collaborator.id, work_date=date(2026, 8, 1)
    )
    assert _record_collaborator_values(record, collaborator)["tax_code"] == CF154


@pytest.mark.parametrize("value", ["invalid", "PLLGNN63B20E972A", "   "])
def test_rejects_invalid_tax_codes(value):
    from app.modules.presenze.services.tax_code_import import normalize_tax_code

    with pytest.raises(ValueError):
        normalize_tax_code(value)


def test_extraction_handles_empty_sheets_and_numeric_record_keys(tmp_path):
    workbook = Workbook()
    path = tmp_path / "empty.xlsx"
    workbook.save(path)
    assert extract_tax_codes(path) == []
    worksheet = workbook.create_sheet("Riepilogo")
    worksheet.append(["Matricola", "Nome", "Record"])
    worksheet.append([154, "", "8/2026-154"])
    worksheet.append([2176, "", None])
    worksheet.append([155, "", CF155.lower()])
    workbook.save(path)
    assert extract_tax_codes(path) == [TaxCodeImportEntry("155", CF155, "Riepilogo!C4")]


def test_import_validates_and_deduplicates_before_querying():
    db = _FakeDb([])
    with pytest.raises(ValueError):
        import_tax_codes(db, [TaxCodeImportEntry("154", "invalid", "test")], apply=True)
    with pytest.raises(ValueError):
        import_tax_codes(db, [TaxCodeImportEntry("", CF154, "test")], apply=True)
    with pytest.raises(ValueError):
        import_tax_codes(db, [TaxCodeImportEntry("154", "", "test")], apply=True)
    with pytest.raises(ValueError, match="discordanti"):
        import_tax_codes(
            db,
            [TaxCodeImportEntry("154", CF154, "a"), TaxCodeImportEntry("154", CF155, "b")],
            apply=True,
        )
    collaborator = SimpleNamespace(employee_code="154", tax_code=None)
    db = _FakeDb([collaborator])
    entries = [TaxCodeImportEntry("154", CF154.lower(), "a"), TaxCodeImportEntry("154", CF154, "b")]
    assert import_tax_codes(db, entries, apply=True, company_code="53")["applied"] == 1
    assert collaborator.tax_code == CF154
    assert import_tax_codes(db, entries, apply=True)["unchanged"] == 1


def test_ambiguous_employee_blocks_import():
    db = _FakeDb([SimpleNamespace(employee_code="154", tax_code=None)] * 2)
    with pytest.raises(TaxCodeImportError) as error:
        import_tax_codes(db, [TaxCodeImportEntry("154", CF154, "test")], apply=True)
    assert error.value.report["conflicts"] == 1


def test_real_transaction_and_audit_failure_leave_database_unchanged():
    from sqlalchemy import create_engine, select
    from sqlalchemy.orm import Session

    from app.db.base import Base

    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        person = PresenzeCollaborator(employee_code="154", company_code="53", name="Test")
        other_company = PresenzeCollaborator(employee_code="154", company_code="54", name="Test")
        db.add_all([person, other_company])
        db.commit()
        entries = [TaxCodeImportEntry("154", CF154, "test")]

        def broken_audit(report, changes):
            assert report["changed"] == 1
            assert changes == [(person, CF154)]
            raise OSError("audit unavailable")

        with pytest.raises(OSError):
            import_tax_codes(db, entries, apply=True, company_code="53", before_commit=broken_audit)
        db.expire_all()
        assert person.tax_code is None
        report = import_tax_codes(
            db, entries, apply=True, company_code="53", before_commit=lambda *_: None
        )
        assert report["applied"] == 1
        db.expire_all()
        assert (
            db.scalar(
                select(PresenzeCollaborator).where(PresenzeCollaborator.id == person.id)
            ).tax_code
            == CF154
        )
        assert other_company.tax_code is None
        assert import_tax_codes(db, entries, apply=False, company_code="53")["unchanged"] == 1
    engine.dispose()


def test_workbook_metadata_recovers_missing_tax_codes():
    from app.modules.presenze.services.xlsm_metadata import load_workbook_metadata

    workbook = Workbook()
    assert load_workbook_metadata(workbook) == {}
    sheet = workbook.create_sheet("Operai")
    sheet.cell(2, 4, 154)
    sheet.cell(2, 16, CF154)
    summary = workbook.create_sheet("Riepilogo")
    summary.append(["Matricola", "Nome", "Record"])
    summary.append([154, "", CF154])
    summary.append([155, "", CF155])
    metadata = load_workbook_metadata(workbook)
    assert metadata[154].tax_code == CF154
    assert metadata[155].tax_code == CF155


def test_cli_selection_manifest_and_atomic_apply(tmp_path, monkeypatch):
    import importlib.util
    import json
    import sys

    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session

    from app.db.base import Base

    script = Path(__file__).parents[1] / "scripts/import_presenze_tax_codes.py"
    spec = importlib.util.spec_from_file_location("import_presenze_tax_codes", script)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    db_url = f"sqlite:///{tmp_path / 'cli.db'}"
    engine = create_engine(db_url)
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        db.add(PresenzeCollaborator(employee_code="154", company_code="53", name="Test"))
        db.commit()
    monkeypatch.setattr(module, "SessionLocal", lambda: Session(engine))
    workbook = tmp_path / "source.xlsx"
    _workbook(workbook)
    manifest = tmp_path / "audit.json"
    base = [str(script), str(workbook)]
    for arguments in [
        [str(script), str(tmp_path / "missing.xlsx")],
        [*base, "--apply"],
        [*base, "--apply", "--employee-code", "154"],
        [*base, "--employee-code", "999"],
    ]:
        monkeypatch.setattr(sys, "argv", arguments)
        with pytest.raises(SystemExit) as error:
            module.main()
        assert error.value.code == 2
    monkeypatch.setattr(sys, "argv", base)
    assert module.main() == 1  # 155 is missing: dry run changes nothing.
    monkeypatch.setattr(
        sys, "argv", [*base, "--apply", "--employee-code", "155", "--manifest", str(manifest)]
    )
    assert module.main() == 2
    monkeypatch.setattr(
        sys, "argv", [*base, "--apply", "--employee-code", "154", "--manifest", str(manifest)]
    )
    assert module.main() == 0
    audit = json.loads(manifest.read_text())
    assert audit["report"]["applied"] == 1
    assert audit["changes"][0]["before"] is None
    assert audit["changes"][0]["after"] == CF154
    assert audit["sources"]["154"]["source"] == "Operai!P2"
    monkeypatch.setattr(sys, "argv", [*base, "--employee-code", "154"])
    assert module.main() == 0
    # A fresh interpreter must register the collaborator's user FK before flush.
    import os
    import subprocess

    with Session(engine) as db:
        db.add(PresenzeCollaborator(employee_code="155", company_code="53", name="Test"))
        db.commit()
    result = subprocess.run(
        [
            sys.executable,
            *base,
            "--apply",
            "--employee-code",
            "155",
            "--manifest",
            str(tmp_path / "fresh.json"),
        ],
        env={**os.environ, "DATABASE_URL": db_url},
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout)["applied"] == 1
    engine.dispose()


def test_tax_code_migration_upgrade_and_downgrade():
    import importlib.util

    import sqlalchemy as sa

    from alembic.migration import MigrationContext
    from alembic.operations import Operations

    path = (
        Path(__file__).parents[1]
        / "alembic/versions/20260923_1200_presenze_collaborator_tax_code.py"
    )
    spec = importlib.util.spec_from_file_location("tax_code_migration", path)
    migration = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(migration)
    engine = sa.create_engine("sqlite://")
    with engine.begin() as connection:
        connection.execute(sa.text("CREATE TABLE presenze_collaborators (id INTEGER PRIMARY KEY)"))
        migration.op = Operations(MigrationContext.configure(connection))
        migration.upgrade()
        assert "tax_code" in {
            col["name"] for col in sa.inspect(connection).get_columns("presenze_collaborators")
        }
        assert (
            sa.inspect(connection).get_indexes("presenze_collaborators")[0]["name"]
            == "ix_presenze_collaborators_tax_code"
        )
        migration.downgrade()
        assert "tax_code" not in {
            col["name"] for col in sa.inspect(connection).get_columns("presenze_collaborators")
        }
    engine.dispose()
