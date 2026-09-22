import json
from datetime import date, timedelta
from pathlib import Path
from types import SimpleNamespace
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook

from app.modules.presenze.models import PresenzeCollaborator, PresenzeDailyRecord
from app.modules.presenze.services import xlsm_export
from app.modules.presenze.services.schedule_engine import DayClassification
from app.modules.presenze.services.xlsm_export import (
    ExportTimesheetRow,
    compile_workbook,
    minutes_to_excel_hours,
    minutes_to_excel_hours_minutes,
    resolve_export_trasferta_value,
)


def test_weekly_ordinary_only_acceptance_cases_match_gate(tmp_path: Path):
    cases = json.loads((Path(__file__).parent / "fixtures/paid-rest-days.json").read_text())
    rows = []
    for index, case in enumerate(cases):
        collaborator = PresenzeCollaborator(
            id=uuid4(), employee_code=str(9000 + index), name=case["name"],
            contract_kind=case.get("contract_kind", "operaio"),
        )
        start = date.fromisoformat(case.get("start", "2026-09-07"))
        weeks = [case["weeks"][0] + [0]] if case.get("sunday") else case["weeks"]
        records = []
        for week_index, week in enumerate(weeks):
            for offset, ordinary in enumerate(week):
                if ordinary is None:
                    continue
                work_date = start + timedelta(days=week_index * 7 + offset)
                records.append(PresenzeDailyRecord(
                    id=uuid4(), collaborator_id=collaborator.id, work_date=work_date,
                    ordinary_minutes=ordinary,
                    straordinario_minutes=case.get("extra", {}).get(work_date.isoformat(), 0),
                    justified_minutes=case.get("justified", {}).get(work_date.isoformat(), 0),
                ))
        rows.append(ExportTimesheetRow(collaborator, records, {}))
    template = tmp_path / "template.xlsx"
    output = tmp_path / "result.xlsm"
    workbook = Workbook()
    workbook.active.title = "Archivio"
    workbook.create_sheet("Archivio2")
    workbook.create_sheet("Giornaliera2")
    workbook.save(template)
    workbook.close()
    compile_workbook(template=template, output=output, rows=rows, period_start=date(2026, 9, 1))
    result = load_workbook(output, keep_vba=True)
    try:
        for index, case in enumerate(cases):
            assert result["Archivio"].cell(index + 2, 25).value == case["expected_worked"], case["name"]
            assert result["Archivio"].cell(index + 2, 26).value == case["expected_paid"], case["name"]
        for cell, column in (("C32", "Y"), ("AJ42", "Y"), ("C33", "Z"), ("AJ43", "Z")):
            assert result["Giornaliera2"][cell].value == (
                f"=SUMIF(Archivio!$C$2:$C$21,$B$12,Archivio!${column}$2:${column}$21)"
            )
    finally:
        xlsm_export.close_workbook_resources(result)


@pytest.mark.parametrize(("raw", "classified", "saturday_extra", "expected"), [
    (456, 420, 0, 5), (420, 456, 0, 6), (456, 456, 60, 6),
])
def test_paid_rest_uses_the_same_classification_as_export(
    monkeypatch, raw, classified, saturday_extra, expected,
):
    collaborator = PresenzeCollaborator(
        id=uuid4(), employee_code="9000", name="Classified", contract_kind="operaio",
    )
    records = [PresenzeDailyRecord(
        id=uuid4(), collaborator_id=collaborator.id,
        work_date=date(2026, 9, day), ordinary_minutes=raw if day < 12 else 0,
    ) for day in range(7, 13)]
    context = object()

    def classify(person, daily, punches, supplied_context):
        assert person is collaborator
        assert supplied_context is context
        assert punches == []
        return DayClassification(
            ordinary_minutes=classified if daily.work_date.day < 12 else 0,
            extra_minutes=0 if daily.work_date.day < 12 else saturday_extra,
            special_day=False, holiday_kind=None, grants_recovery_day=False, source="test",
        )

    monkeypatch.setattr(xlsm_export, "classify_daily_record", classify)
    workbook = Workbook()
    try:
        xlsm_export.write_archivio_summary_values(
            workbook.active, 2, ExportTimesheetRow(collaborator, records, {}),
            period_start=date(2026, 9, 1), schedule_context=context,
        )
        assert workbook.active.cell(2, 16).value == classified * 5 / 60
        assert workbook.active.cell(2, 26).value == expected
    finally:
        workbook.close()


@pytest.mark.parametrize(
    ("minutes", "expected"),
    [
        (None, None),
        (0, 0),
        (1, 0.01),
        (5, 0.05),
        (20, 0.20),
        (30, 0.30),
        (40, 0.40),
        (59, 0.59),
        (60, 1),
        (61, 1.01),
        (390, 6.30),
        (1510, 25.10),
    ],
)
def test_duration_formats_preserve_minutes(minutes, expected):
    assert minutes_to_excel_hours_minutes(minutes) == expected
    assert minutes_to_excel_hours(minutes) == (None if minutes is None else minutes / 60)
    record = SimpleNamespace(trasferta_minutes=minutes, trasferta_montano=False)
    assert resolve_export_trasferta_value(record) == expected
    record.trasferta_montano = True
    assert resolve_export_trasferta_value(record) == "X"


def test_workbook_keeps_daily_hm_and_monthly_decimals(tmp_path: Path, monkeypatch):
    template = tmp_path / "template.xlsx"
    output = tmp_path / "output.xlsm"
    workbook = Workbook()
    workbook.active.title = "Archivio"
    workbook.create_sheet("Archivio2")
    view = workbook.create_sheet("Giornaliera2")
    view["E5"] = "=Archivio2!H5"
    view["AQ5"] = "=INT(E5)+(E5-INT(E5))/60%"
    view["AJ5"] = "=SUM(AQ5:BU5)"
    workbook.save(template)
    workbook.close()
    collaborator = PresenzeCollaborator(id=uuid4(), employee_code="9000", name="Test")
    records = [
        PresenzeDailyRecord(
            id=uuid4(),
            collaborator_id=collaborator.id,
            work_date=date(2026, 9, day),
            ordinary_minutes=390,
            straordinario_minutes=40,
            trasferta_minutes=40,
            km_value=6.5,
        )
        for day in (1, 2)
    ]
    monkeypatch.setattr(
        xlsm_export,
        "resolve_day_classification",
        lambda *_: DayClassification(
            special_day=False,
            ordinary_minutes=390,
            extra_minutes=40,
            holiday_kind=None,
            grants_recovery_day=False,
            source="test",
            overtime_day_minutes=40,
        ),
    )
    compile_workbook(
        template=template,
        output=output,
        rows=[ExportTimesheetRow(collaborator, records, {})],
        period_start=date(2026, 9, 1),
    )
    result = load_workbook(output, keep_vba=True)
    try:
        archive, detail, view = result["Archivio"], result["Archivio2"], result["Giornaliera2"]
        for day in (1, 2):
            for offset, expected in [(0, 6.30), (155, 0.40), (498, 0.40), (279, 6.5)]:
                cell = detail.cell(5, 8 + day - 1 + offset)
                assert cell.value == expected
                assert cell.number_format == "0.00"
        assert archive["H2"].value == 13
        assert archive["L2"].value == pytest.approx(80 / 60)
        assert archive["R2"].value == pytest.approx(860 / 60)
        assert archive["X2"].value == pytest.approx(80 / 60)
        assert view["E5"].value == "=Archivio2!H5"
        assert view["AQ5"].value == "=INT(E5)+(E5-INT(E5))/60%"
        assert view["AJ5"].value == "=SUM(AQ5:BU5)"
        assert view["AJ22"].value == (
            "=SUMPRODUCT(IFERROR(INT(E22:AI22),0))"
            "+SUMPRODUCT(IFERROR(ROUND(MOD(E22:AI22,1)*100,0),0))/60"
        )
    finally:
        result.close()
        result.vba_archive.close()


def test_legacy_metadata_fallbacks_and_workbook_without_summary(tmp_path: Path):
    xlsm_export.close_workbook_resources(SimpleNamespace())
    assert xlsm_export.normalize_request_display_label(" CODE - ") == "CODE -"
    assert xlsm_export.normalize_request_display_label(" - Label ") == "- Label"
    assert xlsm_export.normalize_request_prefix(" CODE - ") is None
    assert xlsm_export.normalize_request_prefix(" - Label ") is None
    assert (
        xlsm_export.resolve_export_absence_code(
            SimpleNamespace(
                request_description=None,
                resolved_absence_cause="not-a-legacy-cause",
            )
        )
        is None
    )
    assert (
        xlsm_export.build_archive_record_key(
            "key- ", period_start=date(2026, 9, 1), employee_code="9000"
        )
        == "9/2026-key-"
    )
    template, output = tmp_path / "legacy.xlsx", tmp_path / "legacy.xlsm"
    workbook = Workbook()
    workbook.active.title = "Archivio2"
    workbook.create_sheet("Giornaliera")
    workbook.save(template)
    workbook.close()
    collaborator = PresenzeCollaborator(id=uuid4(), employee_code="9000", name="Test")
    compile_workbook(
        template=template,
        output=output,
        rows=[ExportTimesheetRow(collaborator, [], {})],
        period_start=date(2026, 9, 1),
    )
    result = load_workbook(output, keep_vba=True)
    try:
        assert result["Archivio2"]["B5"].value == 9000
        assert result["Giornaliera"]["A3"].value == 9
    finally:
        result.close()
        result.vba_archive.close()


@pytest.mark.parametrize(
    ("weekday_minutes", "saturday_minutes", "paid_days"),
    [
        (2279, 0, 0),
        (2280, 0, 1),
        (2281, 0, 1),
        (2280, 60, 0),
    ],
)
def test_paid_saturday_threshold_preserves_minute_precision(
    weekday_minutes, saturday_minutes, paid_days
):
    collaborator = PresenzeCollaborator(
        id=uuid4(), employee_code="9000", name="Test", contract_kind="operaio"
    )
    records = [
        PresenzeDailyRecord(
            id=uuid4(),
            collaborator_id=collaborator.id,
            work_date=date(2026, 9, day),
            ordinary_minutes=456 if day < 11 else weekday_minutes - 4 * 456,
        )
        for day in range(7, 12)
    ]
    records.append(
        PresenzeDailyRecord(
            id=uuid4(),
            collaborator_id=collaborator.id,
            work_date=date(2026, 9, 12),
            ordinary_minutes=saturday_minutes,
        )
    )
    assert (
        xlsm_export.count_operai_paid_rest_days(ExportTimesheetRow(collaborator, records, {}))
        == paid_days
    )
